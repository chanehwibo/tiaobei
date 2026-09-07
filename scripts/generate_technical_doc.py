"""Generate an editable DOCX from the technical solution and test report."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PROJECT_ROOT / "deliverables" / "doc" / "OS_Agent_记忆优化_技术方案与测试报告.docx"


def _plain_markdown(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return text.strip()


def _set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(_plain_markdown(text))
    _set_run_font(run, 9, bold)


def _add_markdown(document: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("<!--") and stripped.endswith("-->"):
            index += 1
            continue
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["No Spacing"]
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(_plain_markdown(heading.group(2)), level=level)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[_plain_markdown(cell) for cell in row.strip("|").split("|")] for row in table_lines]
            column_count = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=column_count)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column_index in range(column_count):
                    _set_cell_text(
                        table.cell(row_index, column_index),
                        row[column_index] if column_index < len(row) else "",
                        bold=row_index == 0,
                    )
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet or numbered:
            value = (bullet or numbered).group(1)
            style = "List Bullet" if bullet else "List Number"
            paragraph = document.add_paragraph(style=style)
            run = paragraph.add_run(_plain_markdown(value))
            _set_run_font(run, 10.5)
            index += 1
            continue
        paragraph = document.add_paragraph()
        if stripped.startswith(">"):
            paragraph.style = "Intense Quote"
            stripped = stripped.lstrip("> ")
        run = paragraph.add_run(_plain_markdown(stripped))
        _set_run_font(run, 10.5)
        paragraph.paragraph_format.line_spacing = 1.35
        index += 1


def build_document(solution_path: Path, report_path: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    run = title.add_run("OS Agent 多源融合记忆优化系统")
    _set_run_font(run, 26, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("技术方案与测试报告（第五轮内部验证版）")
    _set_run_font(run, 16, True)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    evidence = json.loads((PROJECT_ROOT / "artifacts/verification/round5/evidence_summary.json").read_text(encoding="utf-8"))
    run = meta.add_run("对应赛题：XA-202612\n第五轮证据日期：" + evidence["generated_at_utc"][:10])
    _set_run_font(run, 11)
    boundary = document.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    boundary.paragraph_format.space_before = Pt(80)
    run = boundary.add_run("边界：Windows/V11公开合成回归五项达标；原生模型、独立盲测及赛事验收未完成。INCOMPLETE")
    _set_run_font(run, 10, True)

    document.add_page_break()
    _add_markdown(document, solution_path)
    document.add_section(WD_SECTION.NEW_PAGE)
    _add_markdown(document, report_path)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("OS Agent 记忆优化 · 第五轮内部验证 · 结论边界见正文")
        _set_run_font(run, 8)

    document.core_properties.title = "OS Agent 多源融合记忆优化系统：技术方案与测试报告"
    document.core_properties.subject = "挑战杯揭榜挂帅 XA-202612 第五轮内部验证材料"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = "由仓库脚本从 Markdown 正文生成，可编辑；不构成真实麒麟 SDK 验收。"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--solution", type=Path, default=PROJECT_ROOT / "docs" / "09_技术方案正文.md")
    parser.add_argument("--report", type=Path, default=PROJECT_ROOT / "docs" / "测试报告.md")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_document(args.solution, args.report, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
