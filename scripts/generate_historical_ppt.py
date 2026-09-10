"""Reproduce the archived presentation only; never use it as current evidence."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "ppt" / "OS_Agent_记忆优化_答辩稿.pptx"
SUMMARY = ROOT / "artifacts" / "verification" / "round5" / "evidence_summary.json"

NAVY = RGBColor(7, 17, 31)
PANEL = RGBColor(17, 34, 53)
PANEL_2 = RGBColor(24, 48, 72)
WHITE = RGBColor(238, 246, 255)
MUTED = RGBColor(163, 187, 210)
TEAL = RGBColor(69, 222, 196)
BLUE = RGBColor(83, 147, 255)
AMBER = RGBColor(255, 193, 92)
RED = RGBColor(255, 111, 126)
GREEN = RGBColor(99, 222, 139)


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 18,
    color: RGBColor = WHITE,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = "Microsoft YaHei",
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: RGBColor = PANEL,
    line: RGBColor = PANEL_2,
    color: RGBColor = WHITE,
    size: float = 17,
    radius: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = TEAL):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def base_slide(prs: Presentation, title: str, section: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    add_textbox(slide, 0.65, 0.35, 10.8, 0.55, title, size=26, bold=True)
    add_textbox(slide, 11.15, 0.35, 1.5, 0.45, section, size=11, color=TEAL, align=PP_ALIGN.RIGHT)
    add_textbox(
        slide,
        0.65,
        7.16,
        12.0,
        0.22,
        f"XA-202612  ·  OS Agent 记忆优化  ·  {len(prs.slides):02d}",
        size=8,
        color=MUTED,
    )
    return slide


def bullet_list(slide, items: list[str], x: float, y: float, w: float, h: float, *, size: int = 18):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = WHITE
        paragraph.level = 0
        paragraph.space_after = Pt(11)
        paragraph.text = "•  " + paragraph.text
    return shape


def metric_card(slide, x: float, y: float, w: float, label: str, value: str, note: str, color: RGBColor):
    add_box(slide, x, y, w, 1.35, "", fill=PANEL, line=color)
    add_textbox(slide, x + 0.18, y + 0.12, w - 0.36, 0.3, label, size=11, color=MUTED)
    add_textbox(slide, x + 0.18, y + 0.42, w - 0.36, 0.48, value, size=25, color=color, bold=True)
    add_textbox(slide, x + 0.18, y + 0.95, w - 0.36, 0.25, note, size=9, color=MUTED)


def add_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_textbox(slide, 0.9, 1.15, 11.2, 0.45, "挑战杯“揭榜挂帅” · XA-202612", size=17, color=TEAL, bold=True)
    add_textbox(slide, 0.9, 1.85, 11.2, 1.5, "OS Agent 多源融合\n记忆优化与高效应用", size=40, bold=True)
    add_textbox(slide, 0.92, 3.65, 10.6, 0.85, "有作用域 · 有证据 · 有生命周期 · 可冲突治理 · 可安全遗忘", size=20, color=MUTED)
    add_box(slide, 0.92, 5.35, 3.2, 0.7, "Windows / V11 内部回归", fill=PANEL_2, line=TEAL, color=TEAL, size=14)
    add_box(slide, 4.35, 5.35, 3.2, 0.7, "正式提交仍为 INCOMPLETE", fill=PANEL_2, line=AMBER, color=AMBER, size=14)
    evidence = json.loads(SUMMARY.read_text(encoding="utf-8"))
    add_textbox(slide, 0.92, 6.65, 11.4, 0.35, "第五轮逐请求实测稿 · " + evidence["generated_at_utc"][:10], size=12, color=MUTED)


def build_presentation() -> Presentation:
    evidence = json.loads(SUMMARY.read_text(encoding="utf-8"))
    windows = evidence["requests"]["windows"]
    guest = evidence["requests"]["kylin"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "OS Agent 记忆优化：第五轮内部验证"
    prs.core_properties.subject = "XA-202612 公开合成回归；正式提交 INCOMPLETE"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.comments = "从第五轮原始证据生成，不构成独立盲测或官方认证。"
    add_title_slide(prs)

    slide = base_slide(prs, "01  为什么 OS Agent 需要新的记忆系统？", "问题")
    problems = [
        ("偏好失真", "复合偏好被压成单标签；失败工具也可能被学习", RED),
        ("知识冲突", "新旧规则同时返回；矛盾文本被简单拼接", AMBER),
        ("检索退化", "全库向量计算导致端侧时延和资源不可控", BLUE),
        ("安全失控", "敏感信息残留、跨用户串读、模糊指令误删", TEAL),
    ]
    for index, (title, body, color) in enumerate(problems):
        x = 0.72 + (index % 2) * 6.15
        y = 1.25 + (index // 2) * 2.45
        add_box(slide, x, y, 5.65, 1.9, "", fill=PANEL, line=color)
        add_textbox(slide, x + 0.3, y + 0.22, 5.05, 0.45, title, size=21, color=color, bold=True)
        add_textbox(slide, x + 0.3, y + 0.78, 5.05, 0.7, body, size=16, color=WHITE)
    add_textbox(slide, 0.75, 6.25, 11.8, 0.45, "目标：让记忆从“历史堆叠”变成可治理、可评测、可撤回的端侧基础设施", size=19, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "02  赛题要求与交付闭环", "需求")
    requirements = [
        ("多源融合", "6 类事件统一模型", GREEN),
        ("动态偏好", "多维键 + 时间衰减", TEAL),
        ("知识冲突", "默认隔离 + 可追溯决策", AMBER),
        ("≤500ms", "候选召回 + 向量缓存", BLUE),
        ("安全遗忘", "预览 + 确认 + 审计", RED),
        ("短中长期", "TTL + 聚合 + 晋升门控", TEAL),
    ]
    for i, (label, value, color) in enumerate(requirements):
        x = 0.75 + (i % 3) * 4.15
        y = 1.25 + (i // 3) * 2.05
        add_box(slide, x, y, 3.75, 1.55, "", fill=PANEL, line=color)
        add_textbox(slide, x + 0.22, y + 0.2, 3.3, 0.35, label, size=16, color=color, bold=True)
        add_textbox(slide, x + 0.22, y + 0.72, 3.3, 0.45, value, size=15)
    add_box(slide, 0.75, 5.75, 11.95, 0.75, "代码 + SQLite + 服务/API + Web 演示 + 分区评测 + 性能/环境报告", fill=PANEL_2, line=TEAL, color=WHITE, size=17)

    slide = base_slide(prs, "03  系统总体架构", "架构")
    top = ["工具结果", "用户行为", "手动配置", "会话/系统", "导入文档"]
    for i, item in enumerate(top):
        add_box(slide, 0.65 + i * 2.5, 1.12, 2.05, 0.62, item, fill=PANEL_2, line=BLUE, size=13)
    add_line(slide, 6.65, 1.76, 6.65, 2.15)
    add_box(slide, 1.2, 2.15, 10.9, 0.7, "校验 → UTC → 去重 → 递归脱敏 → 记忆污染检测", fill=PANEL, line=TEAL, color=TEAL, size=17)
    add_line(slide, 6.65, 2.87, 6.65, 3.25)
    add_box(slide, 1.2, 3.25, 4.65, 0.78, "多维偏好记忆", fill=PANEL_2, line=GREEN, color=GREEN, size=19)
    add_box(slide, 7.45, 3.25, 4.65, 0.78, "结构化知识记忆", fill=PANEL_2, line=AMBER, color=AMBER, size=19)
    add_line(slide, 3.53, 4.05, 5.1, 4.55)
    add_line(slide, 9.77, 4.05, 8.2, 4.55)
    add_box(slide, 3.85, 4.55, 5.6, 0.78, "短期 TTL → 中期聚合 → 长期门控", fill=PANEL, line=TEAL, size=18)
    add_line(slide, 6.65, 5.35, 6.65, 5.7)
    for i, item in enumerate(["SQLite v2", "冲突治理", "混合检索", "安全遗忘"]):
        add_box(slide, 0.85 + i * 3.12, 5.72, 2.75, 0.76, item, fill=PANEL_2, line=BLUE if i % 2 == 0 else TEAL, size=15)

    slide = base_slide(prs, "04  多维偏好：不再用一个标签概括用户", "偏好")
    keys = ["语言 zh-CN", "详略 brief", "格式 table", "结构 step-by-step", "本地处理", "先脱敏", "高风险确认", "场景工具"]
    for i, key in enumerate(keys):
        x = 0.75 + (i % 4) * 3.08
        y = 1.2 + (i // 4) * 1.08
        add_box(slide, x, y, 2.75, 0.72, key, fill=PANEL, line=TEAL if i < 4 else AMBER, size=14)
    add_textbox(slide, 0.85, 3.65, 5.6, 0.45, "证据评分", size=21, color=TEAL, bold=True)
    add_box(slide, 0.85, 4.2, 5.7, 1.35, "0.30 质量 + 0.35 权威\n+ 0.20 时效 + 0.15 显式性", fill=PANEL_2, line=TEAL, size=19)
    add_textbox(slide, 7.0, 3.65, 5.2, 0.45, "决策机制", size=21, color=AMBER, bold=True)
    bullet_list(slide, ["30 天半衰期，按事件时间而非输入顺序", "失败/取消工具不得成为首选", "同值聚合；变更生成新版本并闭合有效期"], 7.0, 4.1, 5.3, 2.0, size=15)

    slide = base_slide(prs, "05  短期—中期—长期：拒绝一次性噪声", "生命周期")
    stages = [
        ("短期", "会话作用域\n默认 TTL 30 分钟", BLUE),
        ("中期", "主题聚合\n次数 / 来源 / 质量 / 权威", AMBER),
        ("长期", "晋升 / 拒绝\n过期 / 降级", GREEN),
    ]
    for i, (name, body, color) in enumerate(stages):
        x = 0.85 + i * 4.2
        add_box(slide, x, 1.45, 3.5, 1.75, "", fill=PANEL, line=color)
        add_textbox(slide, x + 0.25, 1.68, 3.0, 0.42, name, size=23, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.25, 2.2, 3.0, 0.7, body, size=15, align=PP_ALIGN.CENTER)
        if i < 2:
            add_line(slide, x + 3.55, 2.3, x + 4.08, 2.3)
    add_box(slide, 1.1, 4.1, 11.1, 0.9, "晋升 = 0.30 质量 + 0.30 权威 + 0.25 重复 + 0.15 来源多样性", fill=PANEL_2, line=TEAL, color=WHITE, size=19)
    add_box(slide, 2.15, 5.45, 9.0, 0.75, "风险标记 → 分数 × 0.35 且拒绝晋升；每次决策保留可解释原因", fill=PANEL, line=RED, color=RED, size=16)

    slide = base_slide(prs, "06  知识冲突：先隔离，再决策", "冲突")
    add_box(slide, 0.8, 1.25, 3.25, 1.15, "旧规则\n允许先审批后补票", fill=PANEL, line=AMBER, color=AMBER, size=17)
    add_box(slide, 9.25, 1.25, 3.25, 1.15, "新规则\n禁止先审批后补票", fill=PANEL, line=RED, color=RED, size=17)
    add_line(slide, 4.08, 1.82, 5.2, 1.82, RED)
    add_line(slide, 9.22, 1.82, 8.1, 1.82, RED)
    add_box(slide, 5.15, 1.27, 3.0, 1.1, "conflict_pending\n普通检索不可见", fill=PANEL_2, line=RED, color=WHITE, size=16)
    add_textbox(slide, 0.85, 3.15, 4.2, 0.45, "自动决策优先级", size=20, color=TEAL, bold=True)
    add_box(slide, 0.85, 3.72, 4.25, 1.0, "authority → effective_time → version", fill=PANEL, line=TEAL, size=17)
    add_textbox(slide, 5.65, 3.15, 6.4, 0.45, "支持策略", size=20, color=AMBER, bold=True)
    bullet_list(slide, ["接受新版本 / 保留旧版本 / 保留双方", "证据合并（不拼接矛盾文本）", "差异不足进入人工复核队列"], 5.65, 3.65, 6.2, 1.8, size=15)
    add_box(slide, 1.7, 5.82, 9.9, 0.68, "核心原则：未决知识宁可不回答，也不把冲突建议交给 OS Agent 执行", fill=PANEL_2, line=RED, color=WHITE, size=17)

    slide = base_slide(prs, "07  高效检索：候选召回 + 向量缓存 + 无答案", "检索")
    steps = ["作用域过滤", "多路候选", "查询向量 ×1", "多粒度重排", "阈值/空答案"]
    for i, step in enumerate(steps):
        x = 0.55 + i * 2.58
        add_box(slide, x, 1.55, 2.2, 0.82, step, fill=PANEL, line=TEAL if i % 2 else BLUE, size=14)
        if i < len(steps) - 1:
            add_line(slide, x + 2.2, 1.96, x + 2.52, 1.96)
    add_box(slide, 1.0, 3.05, 11.35, 0.92, "中文：0.25 向量 + 0.25 词法 + 0.45 汉字覆盖 + 0.05 权威度", fill=PANEL_2, line=TEAL, size=20)
    metric_card(slide, 1.05, 4.65, 3.35, "Windows / 10,000 条 P95", f"{windows['metrics']['retrieval_p95_ms']:.3f} ms", "80 个不同请求 / 实际服务", BLUE)
    metric_card(slide, 4.98, 4.65, 3.35, "麒麟 V11 / 10,000 条 P95", f"{guest['metrics']['retrieval_p95_ms']:.3f} ms", "VMware / HashEmbedding", TEAL)
    metric_card(slide, 8.91, 4.65, 3.35, "固定配置", "256 维 / top 3", "阈值 0.265 / 不预热", AMBER)
    add_textbox(slide, 1.0, 6.38, 11.35, 0.34, "进程内 MemoryService.search，非 HTTP / Agent 推理；不是原生 SDK 性能", size=13, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "08  安全遗忘：可预览、可确认、可审计", "安全")
    flow = ["自然语言指令", "作用域解析", "影响预览", "确认令牌", "删除+索引同步", "审计回执"]
    for i, step in enumerate(flow):
        x = 0.42 + i * 2.14
        add_box(slide, x, 1.35, 1.82, 0.78, step, fill=PANEL, line=TEAL if i not in (3, 4) else RED, size=12)
        if i < len(flow) - 1:
            add_line(slide, x + 1.82, 1.74, x + 2.08, 1.74)
    protections = [
        ("跨用户误删", "tenant + user 强制绑定", GREEN),
        ("未确认删除", "预览阶段零修改", BLUE),
        ("残留向量", "正文、状态、倒排索引、向量同步", TEAL),
        ("审计泄露", "不记录原指令和原始正文", AMBER),
    ]
    for i, (risk, control, color) in enumerate(protections):
        x = 0.75 + (i % 2) * 6.15
        y = 3.05 + (i // 2) * 1.42
        add_box(slide, x, y, 5.65, 1.05, "", fill=PANEL_2, line=color)
        add_textbox(slide, x + 0.22, y + 0.16, 1.55, 0.32, risk, size=14, color=color, bold=True)
        add_textbox(slide, x + 1.72, y + 0.16, 3.65, 0.56, control, size=14)
    add_textbox(slide, 0.85, 6.15, 11.6, 0.42, "评测规则：授权集合之外多删任何目标，都计入误删率", size=18, color=RED, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "09  可运行交付：核心、服务、SQLite 与本地仪表盘", "工程")
    layers = [
        ("核心算法", "ingestion / preference / knowledge / lifecycle", BLUE),
        ("应用服务", "ingest · search · conflicts · forget · stats", TEAL),
        ("持久化", "SQLite v2 + v1 迁移 + 向量 + 审计", AMBER),
        ("演示入口", "localhost Web UI + JSON API + 一键预检", GREEN),
    ]
    for i, (name, body, color) in enumerate(layers):
        y = 1.18 + i * 1.32
        add_box(slide, 0.85, y, 2.35, 0.86, name, fill=PANEL_2, line=color, color=color, size=17)
        add_box(slide, 3.48, y, 8.8, 0.86, body, fill=PANEL, line=PANEL_2, size=15)
    add_box(slide, 1.45, 6.08, 10.45, 0.58, "默认监听 127.0.0.1；当前界面用于本地演示，不冒充生产级服务", fill=PANEL_2, line=AMBER, color=AMBER, size=14)

    slide = base_slide(prs, "10  冻结逐请求五项：内部达标，非官方认证", "评测")
    rows = [("偏好准确率 ≥85%", "preference_accuracy", "68 请求；完整键值集合匹配"),
            ("知识召回率 ≥85%", "knowledge_recall", "60 正例；另查无答案/跨域"),
            ("冲突正确率 ≥88%", "conflict_accuracy", "60 对；检测和最终状态联合"),
            ("检索P95 ≤500ms", "retrieval_p95_ms", "80 查询；10,000 条 / Hash"),
            ("敏感F1 ≥90%", "sensitive_f1", "50 请求；实体 TP30/FP0/FN0")]
    for i, (label, key, note) in enumerate(rows):
        y = 1.18 + i * .87
        value = windows["metrics"][key]
        formatted = f"{value:.3f} ms" if key == "retrieval_p95_ms" else f"{value:.2%}"
        add_box(slide, .75, y, 3.7, .63, label, size=16)
        add_textbox(slide, 4.7, y, 2.05, .63, formatted, size=22, color=AMBER, bold=True)
        add_textbox(slide, 7.0, y, 5.3, .63, note, size=15, color=MUTED)
    add_box(slide, .8, 5.7, 11.7, .75, "上表 Windows；V11 同套五项也通过，原始报告分别保存", line=TEAL, color=TEAL, size=17)
    add_textbox(slide, .8, 6.5, 11.7, .35, "公开合成标签由开发者持有；保留失败基线；不等于第三方盲测、真实Agent或官方认证", size=12, color=MUTED)

    slide = base_slide(prs, "11  自动化与安全证据", "验证")
    metric_card(slide, 0.8, 1.22, 3.2, "Windows 自动化测试", str(evidence['tests']), "通过状态与源码摘要见报告", GREEN)
    metric_card(slide, 4.45, 1.22, 3.2, "V11 客户机自动化测试", str(evidence["test_results"]["kylin"]["tests_run"]), "磁盘安装 / VMware / 离线", BLUE)
    metric_card(slide, 8.1, 1.22, 3.2, "真实 SDK 单探针", "768 维", "身份/质量/性能仍待验证", AMBER)
    tests = ["跨租户/用户串读与误删", "提示注入与记忆污染", "失败工具偏好污染", "未决冲突普通检索隔离", "硬删除 SQLite 重载无残留", "嵌套凭证递归脱敏"]
    for i, item in enumerate(tests):
        x = 0.75 + (i % 2) * 6.15
        y = 3.15 + (i // 2) * 0.92
        add_box(slide, x, y, 5.65, 0.64, "✓  " + item, fill=PANEL, line=TEAL, color=WHITE, size=14)

    slide = base_slide(prs, "12  现场演示路径（约 2 分钟）", "演示")
    demo_steps = [
        ("1", "预检", "区分本地通过与外部未验收"),
        ("2", "多源接入", "展示脱敏和偏好证据"),
        ("3", "知识检索", "场景过滤与无答案"),
        ("4", "冲突治理", "待决隔离后自动/人工策略"),
        ("5", "安全遗忘", "预览影响范围，再确认删除"),
        ("6", "评测证据", f"五项逐请求、{evidence['tests']} 测试、V11 报告"),
    ]
    for i, (number, title, body) in enumerate(demo_steps):
        y = 1.12 + i * 0.91
        add_box(slide, 0.85, y, 0.62, 0.62, number, fill=TEAL, line=TEAL, color=NAVY, size=18)
        add_textbox(slide, 1.72, y, 2.0, 0.62, title, size=17, color=TEAL, bold=True)
        add_textbox(slide, 3.72, y, 8.45, 0.62, body, size=15)
    add_box(slide, 3.2, 6.45, 6.95, 0.4, "一键脚本：scripts/run_demo.ps1", fill=PANEL_2, line=BLUE, color=BLUE, size=13)

    slide = base_slide(prs, "13  证据边界与最后一公里", "计划")
    completed = ["生命周期/显式授权门控", "增量事务与失败回滚", "服务/API/场景演示", "冻结逐请求五项回归", "V11磁盘安装与客户机实测", "旧ABI真实向量单探针"]
    pending = ["原生模型身份/向量库往返", "V11虚拟机赛事认可", "真实Agent与独立LLM Judge", "知情同意案例", "最终视频（不超过7分钟）", "平台/报名/邮件回执"]
    add_textbox(slide, 0.8, 1.15, 5.65, 0.45, "已完成（可复现内部证据）", size=20, color=GREEN, bold=True)
    add_textbox(slide, 6.95, 1.15, 5.55, 0.45, "继续验证与补齐", size=20, color=AMBER, bold=True)
    for i, item in enumerate(completed):
        add_box(slide, 0.8, 1.85 + i * 0.72, 5.65, 0.53, "✓  " + item, fill=PANEL, line=GREEN, color=WHITE, size=14)
    for i, item in enumerate(pending):
        add_box(slide, 6.95, 1.85 + i * 0.72, 5.55, 0.53, "!  " + item, fill=PANEL, line=AMBER, color=WHITE, size=14)
    add_textbox(slide, 0.85, 6.42, 11.6, 0.34, "当前准确表述：Windows/V11 内部回归达标；外部赛事验收尚未完成", size=17, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "14  结论", "总结")
    add_textbox(slide, 1.05, 1.25, 11.25, 0.75, "从“记住更多”走向“记得正确、找得快速、忘得安全”", size=29, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    conclusions = [
        ("正确", "多维偏好 + 权威/时效证据 + 冲突隔离", TEAL),
        ("高效", f"候选合并 + 向量缓存；V11 P95 {guest['metrics']['retrieval_p95_ms']:.3f}ms（Hash）", BLUE),
        ("安全", "递归脱敏 + 作用域 + 两阶段遗忘 + 隐私审计", AMBER),
    ]
    for i, (name, body, color) in enumerate(conclusions):
        x = 0.78 + i * 4.18
        add_box(slide, x, 2.65, 3.75, 2.0, "", fill=PANEL, line=color)
        add_textbox(slide, x + 0.25, 2.98, 3.25, 0.5, name, size=24, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.32, 3.62, 3.1, 0.7, body, size=14, align=PP_ALIGN.CENTER)
    add_box(slide, 2.15, 5.35, 9.0, 0.78, "下一步：原生模型复测 × 真实Agent × 独立盲测 × 授权案例", fill=PANEL_2, line=TEAL, color=TEAL, size=18)
    add_textbox(slide, 1.1, 6.42, 11.15, 0.4, "谢谢各位老师，欢迎提问", size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    return prs


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise ValueError("refusing to overwrite historical material")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation = build_presentation()
    presentation.save(args.output)
    print(f"generated {len(presentation.slides)} historical slides: {args.output}")


if __name__ == "__main__":
    main()
