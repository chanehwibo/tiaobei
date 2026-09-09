"""Build fresh editable supplements from matching Windows/V11 raw evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from os_agent_memory.action_judging import controller_source_identity
from os_agent_memory.experiments import assess_decay_diagnostics
from os_agent_memory.official_contract import environment_identity
from os_agent_memory.request_contract import FROZEN_DATASET_SHA256, assess_frozen_requests
from scripts.check_deliverables import inspect_ooxml
from scripts.material_status import INITIAL_STATUS, SDK_RISK_STATUS
from scripts.verify_kylin_guest import CORE_CHECKS, file_hashes, guest_passed, installed_root
from scripts.verify_release import EXTERNAL_BLOCKERS

EVIDENCE = Path("artifacts/verification/takeover")
OFFICE_NAMES = ("OS_Agent_current.pptx", "OS_Agent_current.docx")
WINDOWS_CHECKS = {
    "compile", "unit_tests", "source_syntax_and_privacy", "frozen_request_gate",
    "frozen_dataset_bytes", "source_unchanged_during_checks", "request_source_snapshot",
    "evaluation_protocol", "internal_quality_85", "benchmark_500ms",
    "service_workload_integrity", "decay_diagnostics", "action_framework",
    "local_preflight", "materials", "editable_materials", "privacy_scan",
}
METRICS = (
    ("preference_accuracy", "偏好逐请求准确率", ">=85%"),
    ("knowledge_recall", "知识 Recall@3", ">=85%"),
    ("conflict_accuracy", "冲突正确率", ">=88%"),
    ("retrieval_p95_ms", "检索 P95", "<=500ms"),
    ("sensitive_f1", "敏感实体 F1", ">=90%"),
)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_service(report):
    rows = report.get("results", [])
    require(report.get("integrity_passed") is True and report.get("rounds") == 30,
            "service workload must pass 30 rounds")
    require([row.get("records") for row in rows] == [100, 1000], "service sizes changed")
    expected = {"sequential": 30, "concurrent": 30, "http": 10}
    for row in rows:
        require(row.get("query_counts") == expected and row.get("correct_queries") == expected,
                "service target counts do not match")
        checks = row.get("integrity_checks", {})
        require(set(checks) == {"sequential_target_recalled", "concurrent_target_recalled",
                               "http_target_recalled", "restart_record_count"}
                and all(value is True for value in checks.values())
                and row.get("integrity_passed") is True
                and row.get("persisted_knowledge") == row["records"] + 4,
                "service integrity details failed")
        for name in ("mixed_query_p95_ms", "http_query_p99_ms"):
            value = row.get(name)
            require(type(value) in (int, float) and math.isfinite(value) and value >= 0,
                    "invalid service timing")
        resources = row.get("process_resources", {})
        require(resources.get("available") is True
                and all(type(resources.get(key)) is int and resources[key] > 0
                        for key in ("rss_after_bytes", "lifetime_peak_rss_bytes")),
                "process RSS evidence is unavailable")


def build_summary(root=ROOT):
    root = Path(root).resolve()
    evidence_hashes = {}

    def read(relative, *, text=False):
        path = (root / relative).resolve()
        require(path.is_relative_to(root) and path.is_file(), "missing or unsafe evidence")
        payload = path.read_bytes()
        evidence_hashes[path.relative_to(root).as_posix()] = hashlib.sha256(payload).hexdigest()
        return payload.decode("utf-8") if text else json.loads(payload)

    current = file_hashes(root)
    require(digest(root / "data/eval/request_regression_v1.json") == FROZEN_DATASET_SHA256,
            "frozen dataset bytes changed")
    controller = controller_source_identity(root)["sha256"]
    windows = read(EVIDENCE / "windows/final_verification.json")
    kylin = read(EVIDENCE / "kylin/kylin_vm.json")
    checks = windows.get("checks", {})
    require(WINDOWS_CHECKS.issubset(checks) and all(value is True for value in checks.values())
            and windows.get("local_release_checks_passed") is True
            and windows.get("competition_submission_ready") is False,
            "Windows engineering gate failed or scope overstated")
    identity = environment_identity(
        kylin.get("system"), kylin.get("compatibility", {}).get("operating_system", {}).get("os_release", {})
    )
    require(guest_passed(kylin) and identity["is_required_v11"]
            and installed_root(kylin.get("root_filesystem", ""))
            and kylin.get("tests_passed") is True and kylin.get("official_validation_ready") is False
            and all(type(kylin.get("checks", {}).get(name, {}).get("returncode")) is int
                    and kylin["checks"][name]["returncode"] == 0 for name in CORE_CHECKS),
            "installed V11 core checks failed")
    tests = {}
    requests = {}
    services = {}
    decay = {}
    readiness = {}
    actions = {}
    for name, report in (("windows", windows), ("kylin", kylin)):
        snapshot = report.get("file_hashes", {})
        unchanged = checks.get("source_unchanged_during_checks") if name == "windows" else report.get("source_unchanged_during_checks")
        require(unchanged is True
                and snapshot.get("lf_sha256") == current["lf_sha256"]
                and snapshot.get("excluded_generated_outputs") == current["excluded_generated_outputs"],
                name + " measured source does not match current source")
        test_count = windows["unit_tests"]["tests_run"] if name == "windows" else kylin["tests_run"]
        test_log = (windows["unit_tests"]["output_tail"] if name == "windows"
                    else read(EVIDENCE / "kylin/unit_tests.log", text=True))
        match = re.search(r"Ran (\d+) tests?", test_log)
        require(type(test_count) is int and test_count > 0 and match is not None
                and int(match.group(1)) == test_count and re.search(r"(?m)^OK(?:\s|$)", test_log)
                and (name != "windows" or windows["unit_tests"].get("passed") is True),
                name + " test count or completion log mismatch")
        tests[name] = {"tests_run": test_count, "passed": True,
                       "engineering_checks": len(checks) if name == "windows" else len(CORE_CHECKS)}
        request = read(EVIDENCE / name / "requests.json")
        require(assess_frozen_requests(request)["passed"], name + " frozen request details failed")
        source = {key: value for key, value in snapshot.get("raw_sha256", {}).items()
                  if key.startswith("src/") and key.endswith(".py")}
        require(bool(source) and request.get("source_sha256") == source, name + " request source mismatch")
        requests[name] = {key: request[key] for key in ("metrics", "observation_counts", "dataset")}
        service = read(EVIDENCE / name / "service_workload.json")
        validate_service(service)
        services[name] = service
        ablations = read(EVIDENCE / name / "ablations.json")
        gate = assess_decay_diagnostics(ablations)
        require(gate["passed"] and gate["comparison"]["case_count"] == 6,
                name + " six-case decay diagnostic failed")
        decay[name] = gate["comparison"]
        ready = read(EVIDENCE / name / "action-readiness.json")
        require(ready.get("controller_source_sha256") == controller
                and ready.get("competition_submission_ready") is False
                and ready.get("agent_executed") is False
                and ready.get("independent_judge_executed") is False,
                name + " readiness source or scope mismatch")
        readiness[name] = ready
        action = read(EVIDENCE / name / "actions.json")
        protocol = action.get("protocol", {})
        require(action.get("execution_completed") is True and action.get("official_accuracy_verified") is False
                and action.get("posthoc_mode") == "fixture_exact"
                and action.get("sessions") == 5 and action.get("turns") == 50
                and protocol.get("cross_session_verified") is True
                and protocol.get("probe_files") == 12 and protocol.get("slot_checks") == 50
                and action.get("posthoc", {}).get("passed") is True,
                name + " cross-session fixture failed or scope overstated")
        actions[name] = {"samples": 1, "sessions": 5, "turns": 50,
                         "probe_files": 12, "action_slots": 50, "independent_judge_executed": False}
    acceptance = read(Path("artifacts/verification/x01-x02/organizer_acceptance_reported_20260909.json"))
    require(acceptance.get("reported_organizer_acceptance") is True
            and acceptance.get("organizer_acceptance_verified") is False,
            "organizer acceptance attribution changed; review the external original separately")
    return {
        "schema_version": 1, "generated_at_utc": max(windows["generated_at_utc"], kylin["generated_at_utc"]),
        "scope": "current_matching_source_public_synthetic_regression_not_official_certification",
        "competition_ready": False, "package_status": INITIAL_STATUS,
        "test_results": tests, "requests": requests, "services": services, "decay": decay,
        "actions": actions, "readiness": readiness, "organizer_acceptance": acceptance,
        "v11": {"installed": True, "required_v11_verified": True, "system": kylin["system"],
                "python": kylin["python"], "machine": kylin["machine"],
                "root_filesystem": kylin["root_filesystem"], "virtualization": kylin["virtualization"],
                "logical_cpu_count": kylin["compatibility"]["resources"]["logical_cpu_count"],
                "physical_memory_gb": kylin["compatibility"]["resources"]["physical_memory_gb"]},
        "source_equivalence": {
            "current_lf_matches_both": True, "files": len(current["lf_sha256"]),
            "windows_kylin_raw_equal": windows["file_hashes"]["raw_sha256"] == kylin["file_hashes"]["raw_sha256"],
            "excluded_generated_outputs": current["excluded_generated_outputs"],
        },
        "evidence_sha256": evidence_hashes, "external_blockers": list(EXTERNAL_BLOCKERS),
    }


def metric_rows(summary):
    rows = []
    for key, label, target in METRICS:
        values = [summary["requests"][name]["metrics"][key] for name in ("windows", "kylin")]
        formatted = [f"{value:.3f}ms" if key == "retrieval_p95_ms" else f"{value:.2%}" for value in values]
        rows.append([label, target, *formatted, str(summary["requests"]["windows"]["observation_counts"][key])])
    return rows


def sections(summary):
    win = summary["test_results"]["windows"]
    vm = summary["test_results"]["kylin"]
    source = summary["source_equivalence"]
    decay = summary["decay"]["windows"]
    service = {name: summary["services"][name]["results"][-1] for name in ("windows", "kylin")}
    return [
        ("OS Agent 多源融合记忆优化系统", [
            "赛题 XA-202612 | 当前双环境验证补充版 | " + summary["generated_at_utc"][:10],
            "公开合成冻结回归，HashEmbedding；不是第三方盲测或官方认证。",
            "V11虚拟机部署、运行资源及评分认可按参赛者确认处理，原始凭证待归档。",
            "状态：" + INITIAL_STATUS + "。原生模型、真实Agent及独立Judge仍待验证。",
        ]),
        ("问题与记忆分层", [
            "六类输入统一携带租户、用户、场景、时间、来源和隐私级别。",
            "偏好：证据聚合、修正和版本历史；显式设置与隐式行为分开。",
            "知识：短期、中期、长期晋升；过期、待确认冲突与正常记忆隔离。",
            "遗忘：预览、明确确认、索引同步和脱敏审计，避免隐式扩大删除范围。",
        ]),
        ("核心处理链与持久化", [
            "接入校验与递归脱敏 -> 生命周期 -> 偏好/知识 -> 混合检索 -> 服务/API。",
            "MemoryService统一入口；SQLite按主键增量落盘，重启后恢复索引。",
            "事务失败恢复内存状态，跨请求去重，作用域不明时拒绝检索。",
            "当前支持单进程串行写；多进程并发写与生产认证不在已验收范围。",
        ]),
        ("进程资源实测", [
            "新增系统RSS、进程生命周期峰值、CPU时间；不再只报告Python分配峰值。",
            *[f"{name}: 1,000条负载RSS末值 {row['process_resources']['rss_after_bytes']} bytes；"
              f"生命周期峰值 {row['process_resources']['lifetime_peak_rss_bytes']} bytes。"
              for name, row in service.items()],
            "统计仅覆盖当前Python进程，不含独立模型服务；生命周期峰值不是该负载专属峰值。",
        ]),
        ("偏好衰减与作用域修正", [
            "隐式证据整体按时间衰减；参考时间按tenant_id和user_id分别确定。",
            "隐式权重 = 原始证据权重 × 2^(-证据年龄天数 / 半衰期天数)。",
            "显式手动设置保持有效，直到被显式修改；风险事件不推进偏好参考时间。",
            "拒绝布尔、非有限或非法半衰期；偏好ID包含租户，防止跨租户混淆。",
        ]),
        ("检索稳定性与兼容迁移", [
            "内部候选保留原始分数精度，以稳定ID打破平分；公开分数才舍入。",
            "只规范化全角拉丁字母、数字和空格；不改变中文标点或存储原文。",
            "持久化向量记录latin-width-v1，旧向量缺版本时重新编码。",
            "dev冻结阈值仍为0.265；不修改冻结标签、配置或分母以满足测试。",
        ]),
        ("冻结请求口径", [
            "10,000条记录；HashEmbedding 256维、ngram=2、top_k=3、阈值0.265。",
            "偏好68、知识60、冲突60、检索延迟80、敏感实体50个不同请求。",
            "实体F1是实体级微平均；延迟实测MemoryService.search，不是缓存报告值。",
            "五项数值、配置、请求明细、聚合一致性和附加安全检查共同决定退出码。",
        ]),
        ("五项公开合成回归结果", [
            "下表只证明固定内部数据与当前配置，不代表真实模型泛化能力或赛事认证。",
        ]),
        ("真实安装V11与源码绑定", [
            f"Windows：{win['tests_run']}项测试，{win['engineering_checks']}项工程检查通过。",
            f"V11：{vm['tests_run']}项测试，{vm['engineering_checks']}项核心检查通过；磁盘安装，非Live系统。",
            f"V11 Python {summary['v11']['python']}；{summary['v11']['logical_cpu_count']}逻辑CPU；"
            f"客户机可见内存 {summary['v11']['physical_memory_gb']} GiB。",
            f"{source['files']}项公开被测文件的LF摘要匹配当前源码；双环境原始字节相同："
            f"{source['windows_kylin_raw_equal']}。仅排除已声明的旧自动生成输出。",
        ]),
        ("真实服务负载与重启", [
            "100和1,000条知识：顺序30、并发30、HTTP10次，均检查目标知识命中。",
            "并发查询期间追加3条；加上单条增量写入，重启后分别恢复104和1,004条。",
            *[f"{name}: 1,000条查询P95 {row['mixed_query_p95_ms']:.3f}ms；"
              f"HTTP P99 {row['http_query_p99_ms']:.3f}ms。" for name, row in service.items()],
            "这些小样本P99是经验统计，不是生产SLA，也不是80条冻结延迟请求的分母。",
        ]),
        ("可区分衰减诊断", [
            f"六项定向诊断：启用衰减{decay['with_decay_passed']}/{decay['case_count']}；"
            f"关闭{decay['without_decay_passed']}/{decay['case_count']}。",
            f"改善{decay['improved']}项、持平{decay['tied']}项、退化{decay['regressed']}项。",
            "覆盖隐式风格反转、工具偏好反转、稳定对照、手动持久/修改和失败工具负对照。",
            "开发者编写且已知标签，不计为独立盲测请求，不外推为真实用户效果。",
        ]),
        ("跨会话框架与真实评测边界", [
            "双环境各跑一个合成样例：5会话50轮、12个实际输出文件、50个行动槽。",
            "当前Judge协议共76个检查；50轮或76个检查均不能改称独立请求数。",
            "已有固定镜像、冻结清单、独立Judge交换、逐槽评分及人工复核导入能力。",
            "只读盘点尚未确认可运行Agent/Judge资源；本轮未启动镜像或执行模型。",
        ]),
        ("安全边界与公开交付", [
            "保留失败证据、历史Office和五个既有封存ZIP，当前版单独生成。",
            "公开允许列表排除密钥、数据库、私有案例、模型运行目录与个人报名资料。",
            "Office结构/文本扫描不等同于视觉复核；图片、视频与真实案例需另行检查。",
            "本地HTTP只绑定回环，未提供生产登录认证/TLS，不应直接公开部署。",
        ]),
        ("复现与交付检查", [
            "python -S -m unittest discover -s tests -v",
            "python -S scripts/check_source.py",
            "python -S scripts/evaluate_requests.py --require-frozen --output artifacts/runtime/new.json",
            "完整发布验收使用verify_release.py；V11使用verify_kylin_guest.py。",
            "原始证据、源码摘要与材料哈希分开记录；最终草案独立解包复测。",
        ]),
        ("外部验收与最终提交", [
            "X01：真实Agent/Judge及资源适配、独立标签保管与真实人工复核。",
            "X02：已认可作为工作前提；归档主办方原始凭证及适用配置。",
            "X03：经授权且匿名化的真实应用案例。",
            "X04：不超过7分钟的最终视频复核、平台审核和盖章报名表。",
            "X05：负责人核准后的正式提交和回执。赛事材料不自动发送；" + INITIAL_STATUS + "。",
        ]),
    ]


def render_markdown(summary):
    lines = ["# 当前双环境技术与测试补充报告", "",
             "本补充报告与当前版PPT、DOCX共用本目录的evidence_summary.json。旧版材料保留为历史。", ""]
    for index, (title, body) in enumerate(sections(summary), 1):
        lines += [f"## {index}. {title}", ""]
        lines.extend("- " + line for line in body)
        lines.append("")
        if index == 8:
            lines += ["| 指标 | 门槛 | Windows | 麒麟V11 | 请求数 |",
                      "| --- | --- | --- | --- | --- |"]
            lines += ["| " + " | ".join(row) + " |" for row in metric_rows(summary)]
            lines.append("")
    lines += ["## 证据索引", "", "冻结数据SHA256：`" + FROZEN_DATASET_SHA256 + "`。", ""]
    lines += ["- `" + path + "`：`" + sha + "`。" for path, sha in summary["evidence_sha256"].items()]
    return "\n".join(lines) + "\n"


def build_office(summary, output):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt as DocPt
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from scripts.generate_technical_doc import _add_markdown

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in (("Normal", 10.5), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = DocPt(size)
    _add_markdown(document, output / "current_report.md")
    footer = section.footer.paragraphs[0]
    footer.text = "OS Agent | 当前双环境内部验证 | " + INITIAL_STATUS
    footer.runs[0].font.size = DocPt(8)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "OS Agent 当前双环境技术与测试补充报告"
    document.save(output / OFFICE_NAMES[1])

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.title = "OS Agent 当前双环境验证答辩补充稿"

    def style_paragraph(paragraph, size, color, bold=False):
        font = paragraph.font
        font.name, font.size, font.bold = "Microsoft YaHei", Pt(size), bold
        font.color.rgb = RGBColor(*color)
        east_asia = OxmlElement("a:ea")
        east_asia.set("typeface", "Microsoft YaHei")
        font._rPr.append(east_asia)

    for index, (title, body) in enumerate(sections(summary), 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        def text_box(x, y, width, height, lines, size, color=(37, 43, 42), bold=False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            frame.word_wrap = True
            for number, text in enumerate(lines):
                paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
                paragraph.text = text
                style_paragraph(paragraph, size, color, bold)
                paragraph.space_after = Pt(18)
            return frame

        text_box(.65, .45, 12, .8, [title], 28, bold=True)
        text_box(.7, 1.5, 11.9, 4.95 if index != 8 else .8, body, 19 if index != 14 else 17)
        if index == 8:
            table = slide.shapes.add_table(6, 5, Inches(.75), Inches(2.6), Inches(11.8), Inches(3.5)).table
            for column, width in zip(table.columns, (3.2, 1.7, 2.4, 2.5, 2.0)):
                column.width = Inches(width)
            rows = [["指标", "门槛", "Windows", "麒麟V11", "请求数"], *metric_rows(summary)]
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(*( (32, 104, 91) if row_index == 0 else (246, 248, 247)))
                    for paragraph in cell.text_frame.paragraphs:
                        style_paragraph(paragraph, 16, (255, 255, 255) if row_index == 0 else (37, 43, 42))
        text_box(.7, 6.92, 10.8, .35, ["公开合成回归 | HashEmbedding | " + INITIAL_STATUS], 11, (149, 51, 53))
        text_box(12, 6.92, .65, .35, [f"{index}/15"], 11)
    presentation.save(output / OFFICE_NAMES[0])


def verify_outputs(root, output, *, require_office=True):
    summary = build_summary(root)
    issues = []
    stored = json.loads((output / "evidence_summary.json").read_text(encoding="utf-8"))
    if stored != summary:
        issues.append("summary_does_not_match_current_evidence")
    if (output / "current_report.md").read_text(encoding="utf-8") != render_markdown(summary):
        issues.append("markdown_does_not_match_evidence")
    reports = {}
    if require_office:
        expected = [INITIAL_STATUS, "V11", "HashEmbedding", "P95"]
        expected += [value for row in metric_rows(summary) for value in row[2:]]
        expected += [str(row["tests_run"]) for row in summary["test_results"].values()]
        for name in OFFICE_NAMES:
            path = output / name
            report = inspect_ooxml(path)
            content = "\n".join(report.pop("texts"))
            report["missing_evidence_values"] = sorted({value for value in expected if value not in content})
            report["passed"] = report["passed"] and not report["missing_evidence_values"]
            report["sha256"] = digest(path) if path.is_file() else None
            reports[name] = report
    return {"passed": not issues and all(row["passed"] for row in reports.values()),
            "scope": "source_bound_evidence_and_ooxml_checks_not_visual_or_official_acceptance",
            "issues": issues, "deliverables": reports, "visual_render_verified": False,
            "competition_ready": False}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "deliverables/current")
    parser.add_argument("--office", action="store_true", help="requires optional python-docx and python-pptx")
    parser.add_argument("--check-only", action="store_true", help="check existing Markdown, summary and both Office files")
    parser.add_argument("--report", type=Path, help="new verification JSON, never overwritten")
    args = parser.parse_args(argv)
    require(args.report is None or not args.report.exists(), "refusing to overwrite material verification")
    if not args.check_only:
        require(not args.output_dir.exists(), "refusing to overwrite current materials; use a fresh directory")
        summary = build_summary(ROOT)
        args.output_dir.mkdir(parents=True, exist_ok=False)
        (args.output_dir / "evidence_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (args.output_dir / "current_report.md").write_text(render_markdown(summary), encoding="utf-8")
        if args.office:
            build_office(summary, args.output_dir)
    report = verify_outputs(ROOT, args.output_dir, require_office=args.check_only or args.office)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        with args.report.open("x", encoding="utf-8") as stream:
            stream.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
