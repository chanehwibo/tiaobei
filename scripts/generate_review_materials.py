"""Generate source-bound review material without claiming current V11 acceptance."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from os_agent_memory.complex_retrieval import FROZEN_SHA256, SUITE_PATH
from os_agent_memory.request_contract import FROZEN_DATASET_SHA256, assess_frozen_requests
from scripts.check_deliverables import inspect_ooxml
from scripts.compare_retrieval_runs import compare_reports, validate_report
from scripts.generate_takeover_materials import METRICS, WINDOWS_CHECKS, digest, require, validate_service
from scripts.material_status import INITIAL_STATUS, SDK_RISK_STATUS
from scripts.summarize_sdk_risk import summarize_sdk_risk
from scripts.verify_kylin_guest import file_hashes

EVIDENCE = Path("artifacts/verification/feedback")
WINDOWS_EVIDENCE = EVIDENCE / "windows-final-03"
OUTPUT = Path("deliverables/review")
OFFICE_NAMES = ("OS_Agent_反馈落实答辩稿.pptx", "OS_Agent_技术方案与复测报告.docx")
DOCUMENTS = (
    "docs/09_技术方案正文.md",
    "docs/38_相关工作与创新边界.md",
    "docs/39_评测口径与可信度.md",
    "docs/40_复杂场景与检索性能复测.md",
)
OBSOLETE = ("INCOMPLETE", "第五轮", "未完成")


def validate_windows_evidence(current, windows, requests, service):
    checks = windows.get("checks", {})
    require(WINDOWS_CHECKS.issubset(checks) and all(value is True for value in checks.values())
            and windows.get("local_release_checks_passed") is True
            and windows.get("competition_submission_ready") is False,
            "Windows engineering checks failed or acceptance was overstated")
    snapshot = windows.get("file_hashes", {})
    require(snapshot.get("lf_sha256") == current["lf_sha256"]
            and snapshot.get("excluded_generated_outputs") == current["excluded_generated_outputs"],
            "Windows measured source does not match current source")
    tests = windows.get("unit_tests", {})
    match = re.search(r"Ran (\d+) tests?", tests.get("output_tail", ""))
    require(tests.get("passed") is True and type(tests.get("tests_run")) is int
            and tests["tests_run"] > 0 and match is not None
            and int(match.group(1)) == tests["tests_run"]
            and re.search(r"(?m)^OK(?:\s|$)", tests["output_tail"]) is not None,
            "Windows test count or completion log mismatch")
    source = {key: value for key, value in snapshot.get("raw_sha256", {}).items()
              if key.startswith("src/") and key.endswith(".py")}
    gate = assess_frozen_requests(requests)
    require(bool(source) and requests.get("source_sha256") == source and gate["passed"]
            and requests.get("environment", {}).get("system") == "Windows"
            and windows.get("request_metrics") == requests["metrics"]
            and windows.get("request_counts") == requests["observation_counts"],
            "Windows frozen request details or source mismatch")
    validate_service(service)
    require(windows.get("service_workload") == service, "Windows service evidence mismatch")
    return gate


def frozen_block(windows):
    lines = [
        "<!-- BEGIN CURRENT WINDOWS EVIDENCE -->",
        "本次当前源码实测仅覆盖Windows开发机的公开合成冻结回归，不是第三方盲测或官方认证。",
        "HashEmbedding 256维、ngram=2、10,000条记录、top_k=3、阈值0.265；实际调用MemoryService.search。",
        "",
        "| 指标 | 门槛 | 当前Windows | 请求数 |",
        "| --- | --- | --- | --- |",
        *["| " + " | ".join(row) + " |" for row in metric_rows({"windows": windows})],
        "",
        f"自动化测试：Windows {windows['tests_run']}项通过；{windows['engineering_checks']}项工程检查通过。",
        "当前源码的V11部署待独立复验，不沿用历史V11通过数据。原生SDK仍是待复测风险项。",
        "五项及附加安全门槛均通过；敏感识别为实体级微平均，不代表任意敏感实体均可识别。",
        "",
        "数据原始SHA256：`" + FROZEN_DATASET_SHA256 + "`。",
        "原始逐请求报告：`" + (WINDOWS_EVIDENCE / "requests.json").as_posix() + "`；",
        "测试日志、检查结果与完整源码摘要：`" + (WINDOWS_EVIDENCE / "final_verification.json").as_posix() + "`。",
        "<!-- END CURRENT WINDOWS EVIDENCE -->",
    ]
    return "\n".join(lines)


def build_summary(root=ROOT):
    root = Path(root).resolve()
    hashes = {}

    def read(relative, *, text=False):
        path = (root / relative).resolve()
        require(path.is_relative_to(root) and path.is_file(), "missing or unsafe review evidence")
        payload = path.read_bytes()
        hashes[path.relative_to(root).as_posix()] = hashlib.sha256(payload).hexdigest()
        return payload.decode("utf-8") if text else json.loads(payload)

    current = file_hashes(root)
    require(digest(root / "data/eval/request_regression_v1.json") == FROZEN_DATASET_SHA256
            and digest(root / SUITE_PATH) == FROZEN_SHA256, "frozen dataset bytes changed")
    windows = read(WINDOWS_EVIDENCE / "final_verification.json")
    requests = read(WINDOWS_EVIDENCE / "requests.json")
    service = read(WINDOWS_EVIDENCE / "service_workload.json")
    validate_windows_evidence(current, windows, requests, service)
    before = read(EVIDENCE / "complex_before.json")
    after = read(EVIDENCE / "complex_after.json")
    suite = read(SUITE_PATH)
    comparison = compare_reports(before, after, suite)
    comparison["input_sha256"] = {
        name: hashes[(EVIDENCE / ("complex_" + name + ".json")).as_posix()]
        for name in ("before", "after")
    }
    require(comparison == read(EVIDENCE / "complex_comparison.json"), "paired comparison was changed")
    require(comparison["performance_improvement_supported"] is True
            and comparison["identical_ranked_results_and_scores"] is True
            and after["safety_passed"] is True and after["latency_target_met"] is True,
            "paired performance or unchanged-output claim is unsupported")
    expected_profiles = {(size, fraction, "hybrid") for size in (1000, 10000) for fraction in (0, .75)}
    require({(row["corpus"]["records"], row["corpus"]["requested_hard_fraction_of_fillers"], row["mode"])
             for row in after["profiles"]} == expected_profiles and after["method"]["rounds"] == 3,
            "paired review needs all four profiles and three repetitions")
    ablations = read(EVIDENCE / "complex_ablations.json")
    validate_report(ablations, suite)
    require({(row["corpus"]["records"], row["corpus"]["requested_hard_fraction_of_fillers"], row["mode"])
             for row in ablations["profiles"]} == {(1000, .75, mode) for mode in ("lexical", "semantic", "hybrid")}
            and ablations["method"]["rounds"] == 1, "ablation profile or repetition changed")
    source = {key: value for key, value in current["raw_sha256"].items()
              if key.startswith("src/") and key.endswith(".py")}
    evaluator_paths = {SUITE_PATH.as_posix(), "src/os_agent_memory/complex_retrieval.py",
                       "scripts/evaluate_complex_retrieval.py"}
    for report in (after, ablations):
        require(report["source_sha256"] == source, "complex measured source does not match current source")
        require(set(report["evaluator_sha256"]) == evaluator_paths
                and all(digest(root / key) == value for key, value in report["evaluator_sha256"].items()),
                "complex evaluator source changed")
    sdk = read(EVIDENCE / "sdk_risk.json")
    require(sdk == summarize_sdk_risk(root), "SDK risk does not match the raw probe")
    read(sdk["source"])
    acceptance = read(Path("artifacts/verification/x01-x02/organizer_acceptance_reported_20260909.json"))
    require(acceptance.get("reported_organizer_acceptance") is True
            and acceptance.get("organizer_acceptance_verified") is False,
            "organizer acceptance attribution changed")
    windows_summary = {
        "tests_run": windows["unit_tests"]["tests_run"], "engineering_checks": len(windows["checks"]),
        "metrics": requests["metrics"], "observation_counts": requests["observation_counts"],
        "environment": requests["environment"], "service": service["results"][-1],
    }
    for relative in DOCUMENTS:
        text = read(relative, text=True)
        require(INITIAL_STATUS in text and not any(value in text for value in OBSOLETE),
                "obsolete status in current document: " + relative)
        require("ROUND5 EVIDENCE" not in text, "replace the historical current-metrics block first")
        if relative == DOCUMENTS[0]:
            require(frozen_block(windows_summary) in text, "solution metrics do not match current Windows evidence")
    profiles = [{
        "records": profile["corpus"]["records"],
        "hard_fraction": profile["corpus"]["requested_hard_fraction_of_fillers"],
        "mode": profile["mode"],
        "rounds": [repeat["summary"] for repeat in profile["rounds"]],
    } for profile in after["profiles"]]
    ablation_rows = [{
        "mode": profile["mode"], **profile["rounds"][0]["summary"],
    } for profile in ablations["profiles"]]
    return {
        "schema_version": 1,
        "generated_at_utc": windows["generated_at_utc"],
        "scope": "current_windows_hash_public_synthetic_review_not_official_acceptance",
        "display_status": INITIAL_STATUS,
        "competition_ready": False,
        "windows": windows_summary,
        "v11_current_source_verified": False,
        "v11_status": "待当前版复验；历史安装或测试结果不替代当前服务部署证明",
        "sdk": sdk,
        "complex": {
            "dataset_sha256": FROZEN_SHA256, "distinct_requests": 120,
            "positive_requests": 72, "no_answer_requests": 24, "scope_requests": 24,
            "profiles": profiles, "comparison": comparison, "ablations": ablation_rows,
            "quality_targets_met": after["quality_targets_met"],
            "safety_passed": after["safety_passed"],
            "latency_target_met": after["latency_target_met"],
        },
        "organizer_acceptance": acceptance,
        "source_lf_sha256": current["lf_sha256"],
        "evidence_sha256": hashes,
    }


def main_complex(summary):
    return next(row["rounds"][0] for row in summary["complex"]["profiles"]
                if row["records"] == 10000 and row["hard_fraction"] == .75)


def metric_rows(summary):
    windows = summary["windows"]
    return [[label, target, f"{windows['metrics'][key]:.3f}ms" if key == "retrieval_p95_ms"
             else f"{windows['metrics'][key]:.2%}", str(windows["observation_counts"][key])]
            for key, label, target in METRICS]


def slide_sections(summary):
    windows, complex_result = summary["windows"], summary["complex"]
    service, quality = windows["service"], main_complex(summary)
    paired = complex_result["comparison"]
    count = sum(row["rounds"][0]["requests"] * len(row["rounds"]) for row in complex_result["profiles"])
    return [
        ("OS Agent 多源融合记忆优化系统", [
            "XA-202612 | 出题方反馈落实与当前源码复测 | " + summary["generated_at_utc"][:10],
            INITIAL_STATUS,
            "端侧记忆治理、作用域隔离、混合检索与可核验的遗忘。",
            "当前成绩：Windows + HashEmbedding + 公开合成数据。",
            "真实SDK：" + SDK_RISK_STATUS + "；V11待当前版复验。",
        ]),
        ("四项反馈的落实范围", [
            "状态：当前材料和新审阅包统一为“" + INITIAL_STATUS + "”。",
            "SDK：保留真实慢调用，明确复测风险和解除条件。",
            "创新与可信度：补五项一手文献对比、分母和时延口径。",
            "复杂干扰：冻结12场景120请求，分层报告质量和延迟。",
            "新增难例揭示近主题误答风险；不能概括为所有复杂质量达标。",
        ]),
        ("方案与工程贡献", [
            "事件 -> 校验/脱敏 -> 授权晋升 -> 偏好/知识 -> 召回/遗忘。",
            "显式设置优先，行为证据衰减；失败工具不能自动成为偏好。",
            "待决冲突、过期、隔离内容不进入普通检索。",
            "tenant/user/scene作用域贯穿索引、持久化、预览和确认。",
            "创新限定为OS Agent约束下的集成一致性，不声称基础算法首创。",
        ]),
        ("相关工作与创新边界", [
            "MemGPT：分层上下文管理；本项目侧重事件授权与治理。",
            "MemoryBank：长期记忆和遗忘；本项目增加显式决策与版本证据。",
            "Generative Agents：观察/反思/规划；本项目不声称复现智能体。",
            "LongMem：训练式记忆读取；本项目为可替换模型的低依赖原型。",
            "FAISS：向量索引；本项目不是ANN创新，也不作跨论文数字比较。",
        ]),
        ("评测口径与性能可信度", [
            "冻结五项分母为68/60/60/80/50；复杂集另报，绝不混分母。",
            "10,000条、Hash256/ngram2、top_k=3、阈值0.265。",
            "Recall取查询宏平均；实体F1取微平均；无答案、越域单列。",
            "P95使用nearest-rank；保留所有异常、超时与逐轮结果。",
            "服务计时包含锁、编码、召回、重排、序列化；不含HTTP/Agent。",
        ]),
        ("当前Windows冻结五项", [
            "公开合成回归；五项数值和附加安全门槛共同校验。",
        ]),
        ("复杂场景与干扰设计", [
            "12个虚构桌面场景，132条固定难例，120个不同查询。",
            "72正例、24近主题无答案、24作用域检查；12例需双证据。",
            "覆盖相似实体、条件差异、版本时效、长文和跨域污染。",
            "1,000/10,000条规模；0/75%相似填充，固定难例始终存在。",
            "75%的分母仅为生成填充记录；数据由开发者编写，不是盲测。",
        ]),
        ("相同结果下的检索延迟优化", [
            f"{count:,}对返回ID、排序和分数一致：{paired['identical_ranked_results_and_scores']}。",
            "下表为各配置首轮P95；完整三轮及首查询耗时随包保留。",
        ]),
        ("优化做了什么", [
            "倒排交集计数，替代逐候选重复构造集合。",
            "有界堆选择候选和top-k，替代全量排序。",
            "明确tenant/user时直接使用授权作用域，保留歧义拒绝。",
            "4096项有界缓存只解析不可变时间字符串，不缓存权限或答案。",
            "仅为最终命中构造结果；独立全排序参考与变更/过期测试兜底。",
        ]),
        ("复杂质量与误召回风险", [
            "10,000条、75%相似填充、首轮；全部请求保留。",
            "作用域零泄漏不等于近主题拒答可靠，后者仍需改进。",
        ]),
        ("本地检索消融", [
            "1,000条、75%相似填充、同一阈值0.265、各一次完整运行。",
            "这是固定阈值诊断，不是调参后的最优基线或论文复现。",
        ]),
        ("真实SDK：待复测风险项", [
            f"旧ABI真实执行1次，768维；总耗时{summary['sdk']['single_probe_ms']:,.3f}ms。",
            "含隔离进程、会话及模型初始化；单次不是P95。",
            "当前provider每次调用初始化，不能假定后续自动摊销。",
            "新ABI缺符号；模型身份/版本及向量库往返尚无完整验证。",
            "解除条件：核验模型，冻结至少50不同请求，全服务P95<=500ms。",
        ]),
        ("工程验证与部署边界", [
            f"Windows：{windows['tests_run']}项测试、{windows['engineering_checks']}项工程检查通过。",
            f"1,000条真实服务负载：查询P95 {service['mixed_query_p95_ms']:.3f}ms；"
            f"HTTP P99 {service['http_query_p99_ms']:.3f}ms。",
            f"进程RSS末值 {service['process_resources']['rss_after_bytes']} bytes；不含模型服务。",
            "V11待当前版复验；历史HTTP曾受系统网络管控影响，失败保留。",
            "主办方部署/资源/评分认可按参赛者确认处理，原件待归档。",
        ]),
        ("复现与证据绑定", [
            "源码、冻结数据、逐请求结果、对比报告分别记录SHA256。",
            "运行unittest、check_source、evaluate_requests、verify_release。",
            "复杂回归使用PYTHONHASHSEED=0；四种规模/干扰配置各三轮。",
            "审阅包采用允许列表；解包后重新校验、运行测试和冻结五项。",
            "旧ZIP和失败记录只作追溯，不覆盖、不冒充当前版通过。",
        ]),
        ("当前结论与验收边界", [
            f"复杂正例Recall@3 {quality['macro_recall_at_3']:.2%}；"
            f"近主题无答案正确率 {quality['no_answer_accuracy']:.2%}。",
            "延迟优化有配对证据；复杂拒答、真实SDK和V11当前版仍有风险。",
            "独立Agent/Judge、保密标签、授权真实案例和赛事材料需外部证据。",
            "技术创新性30分、性能可信度25分仅为反馈组织依据，不是自评得分。",
            INITIAL_STATUS + "；不自动提交或发送邮件。",
        ]),
    ]


def slide_tables(summary):
    quality = main_complex(summary)
    comparison = summary["complex"]["comparison"]
    return {
        6: [["指标", "门槛", "Windows", "请求数"], *metric_rows(summary)],
        8: [["记录数", "相似填充", "优化前P95", "优化后P95"],
            *[[f"{row['records']:,}", f"{row['hard_fraction']:.0%}",
               f"{row['before_p95_ms']:.3f}ms", f"{row['after_p95_ms']:.3f}ms"]
              for row in comparison["rows"] if row["round"] == 1]],
        10: [["质量项", "实测值", "解释"],
             ["正例宏Recall@3", f"{quality['macro_recall_at_3']:.2%}", "不是拒答正确率"],
             ["固定分母P@3", f"{quality['precision_at_3']:.2%}", "按正例查询数乘3"],
             ["近主题无答案", f"{quality['no_answer_accuracy']:.2%}",
              f"{quality['no_answer_correct']}/{quality['no_answer_requests']}正确，{quality['false_answers']}例误答"],
             ["干扰返回占比", f"{quality['distractor_return_rate']:.2%}", "正例返回片段中的无关项"],
             ["跨域泄漏 / 异常", f"{quality['leaked_records']} / {quality['errors']}", "不能覆盖近主题误答风险"]],
        11: [["模式", "Recall@3", "无答案正确率", "P95"],
             *[[row["mode"], f"{row['macro_recall_at_3']:.2%}", f"{row['no_answer_accuracy']:.2%}",
                f"{row['p95_ms']:.3f}ms"] for row in summary["complex"]["ablations"]]],
    }


def render_markdown(summary):
    lines = ["# 当前源码反馈落实与复测摘要", "", "状态：**" + INITIAL_STATUS + "**。", ""]
    tables = slide_tables(summary)
    for index, (title, body) in enumerate(slide_sections(summary), 1):
        lines.extend([f"## {index}. {title}", ""])
        lines.extend("- " + value for value in body)
        lines.append("")
        if index in tables:
            rows = tables[index]
            lines += ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
            lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            lines.append("")
    lines += ["## 可追溯证据", "",
              "冻结五项SHA256：`" + FROZEN_DATASET_SHA256 + "`。",
              "复杂数据SHA256：`" + FROZEN_SHA256 + "`。",
              "完整文件摘要与当前源码LF摘要见同目录evidence_summary.json。", ""]
    return "\n".join(lines)


def build_office(root, summary, output):
    from docx import Document
    from docx.oxml import OxmlElement as DocElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt as DocPt
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from scripts.generate_technical_doc import _add_markdown

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in (("Normal", 10.5), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name, style.font.size = "Microsoft YaHei", DocPt(size)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    _add_markdown(document, output / "current_report.md")
    for relative in DOCUMENTS:
        first_paragraph = len(document.paragraphs)
        _add_markdown(document, Path(root) / relative)
        document.paragraphs[first_paragraph].paragraph_format.page_break_before = True
    for table in document.tables:
        table.rows[0]._tr.get_or_add_trPr().append(DocElement("w:tblHeader"))
        for row in table.rows:
            row._tr.get_or_add_trPr().append(DocElement("w:cantSplit"))
    footer = section.footer.paragraphs[0]
    footer.text = "OS Agent | " + INITIAL_STATUS + " | "
    field = DocElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.runs[0].font.size = DocPt(8)
    document.core_properties.author = document.core_properties.last_modified_by = ""
    document.core_properties.title = "OS Agent 技术方案、相关工作与当前源码复测报告"
    document.core_properties.subject = INITIAL_STATUS
    document.save(output / OFFICE_NAMES[1])

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.author = presentation.core_properties.last_modified_by = ""
    presentation.core_properties.title = "OS Agent 反馈落实答辩稿"
    presentation.core_properties.subject = INITIAL_STATUS
    tables = slide_tables(summary)

    def style_paragraph(paragraph, size, color, bold=False):
        font = paragraph.font
        font.name, font.size, font.bold = "Microsoft YaHei", Pt(size), bold
        font.color.rgb = RGBColor(*color)
        east_asia = OxmlElement("a:ea")
        east_asia.set("typeface", "Microsoft YaHei")
        font._rPr.append(east_asia)

    for index, (title, body) in enumerate(slide_sections(summary), 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        def text_box(x, y, width, height, lines, size, color=(35, 42, 42), bold=False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            frame.word_wrap = True
            for number, text in enumerate(lines):
                paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.space_after = Pt(16)
                style_paragraph(paragraph, size, color, bold)

        text_box(.65, .4, 12, .8, [title], 28, bold=True)
        text_box(.7, 1.45, 11.9, 1.1 if index in tables else 5.1, body,
                 17 if index in tables else 19)
        if index in tables:
            rows = tables[index]
            table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.75), Inches(2.8),
                                           Inches(11.8), Inches(3.6)).table
            widths = (3.9, 2.0, 5.9) if len(rows[0]) == 3 else (3.2, 2.5, 3.5, 2.6)
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    cell.margin_top = cell.margin_bottom = Inches(.08)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(*((32, 104, 91) if row_index == 0 else (244, 247, 246)))
                    for paragraph in cell.text_frame.paragraphs:
                        style_paragraph(paragraph, 16, (255, 255, 255) if row_index == 0 else (35, 42, 42))
        text_box(.7, 6.92, 11, .35, ["公开合成回归 | HashEmbedding | " + INITIAL_STATUS], 11, (145, 51, 54))
        text_box(12, 6.92, .65, .35, [f"{index}/15"], 11)
    presentation.save(output / OFFICE_NAMES[0])


def verify_outputs(root, output, *, require_office=True):
    summary = build_summary(root)
    issues = []
    if json.loads((output / "evidence_summary.json").read_bytes()) != summary:
        issues.append("summary_does_not_match_current_evidence")
    if (output / "current_report.md").read_text(encoding="utf-8") != render_markdown(summary):
        issues.append("markdown_does_not_match_current_evidence")
    reports = {}
    if require_office:
        quality = main_complex(summary)
        expected = [INITIAL_STATUS, SDK_RISK_STATUS, "HashEmbedding", "V11", "MemGPT", "FAISS",
                    "nearest-rank", f"{quality['no_answer_accuracy']:.2%}", f"{quality['macro_recall_at_3']:.2%}",
                    f"{summary['sdk']['single_probe_ms']:,.3f}",
                    str(summary["windows"]["tests_run"])]
        expected += [row[2] for row in metric_rows(summary)]
        for name in OFFICE_NAMES:
            path = output / name
            report = inspect_ooxml(path)
            text = "\n".join(report.pop("texts"))
            report["missing_values"] = sorted({value for value in expected if value not in text})
            report["obsolete_status"] = [value for value in OBSOLETE if value in text]
            report["passed"] = report["passed"] and not report["missing_values"] and not report["obsolete_status"]
            report["sha256"] = digest(path) if path.is_file() else None
            reports[name] = report
    return {
        "passed": not issues and all(row["passed"] for row in reports.values()),
        "scope": "current_windows_source_and_review_evidence_not_current_v11_or_official_acceptance",
        "issues": issues, "deliverables": reports, "visual_render_verified": False,
        "complex_quality_targets_met": summary["complex"]["quality_targets_met"],
        "native_latency_passed": summary["sdk"]["native_latency_passed"],
        "v11_current_source_verified": False, "competition_ready": False,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / OUTPUT)
    parser.add_argument("--office", action="store_true")
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args(argv)
    require(args.report is None or not args.report.exists(), "refusing to overwrite review verification")
    if not args.check_only:
        require(not args.output_dir.exists(), "refusing to overwrite review materials")
        summary = build_summary(ROOT)
        args.output_dir.mkdir(parents=True, exist_ok=False)
        (args.output_dir / "evidence_summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (args.output_dir / "current_report.md").write_text(render_markdown(summary), encoding="utf-8")
        if args.office:
            build_office(ROOT, summary, args.output_dir)
    report = verify_outputs(ROOT, args.output_dir, require_office=args.office or args.check_only)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        with args.report.open("x", encoding="utf-8") as stream:
            stream.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
