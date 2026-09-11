"""Structural checks and private rendering contact sheets for the Office pair."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx import Document
from PIL import Image, ImageDraw, ImageFont, ImageStat
from pptx import Presentation

from docs.materials_20260911.content import STATUS, build_topics
from docs.materials_20260911.diagrams import DIAGRAMS
from docs.materials_20260911.evidence import lf_sha256, load_evidence, load_json, require, sha256
from docs.materials_20260911.generate import (
    DOCX_NAME, PPTX_NAME, layout_source_hashes, presentation_topics, validate_word_layout, write_json,
)
from scripts.check_deliverables import inspect_ooxml

FORBIDDEN = ("INCOMPLETE", "第五轮", "未完成", "CanhuiBao", "chanehwibo", "C:\\Users", "C:/Users")
TEXT_TAGS = {
    "{http://schemas.openxmlformats.org/drawingml/2006/main}t",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
}


def normalized(text):
    return re.sub(r"[\s\u00ad\u200b]+", "", str(text))


def package_text(path):
    values = []
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith((".xml", ".rels")):
                continue
            raw = archive.read(name).decode("utf-8")
            require(not any(token.casefold() in raw.casefold() for token in FORBIDDEN),
                    "identity or obsolete status in Office package: " + name)
            root = ET.fromstring(raw)
            values.extend(node.text or "" for node in root.iter() if node.tag in TEXT_TAGS)
            if name.endswith(".rels"):
                require(not any(node.get("TargetMode") == "External" for node in root),
                        "unexpected external Office relationship")
    return normalized("".join(values))


def verify_render_equivalent(source, rendered_source):
    """Reuse a read-only render only when every uncompressed Office part matches."""
    with ZipFile(source) as actual, ZipFile(rendered_source) as previous:
        names = actual.namelist()
        require(len(names) == len(set(names))
                and len(previous.namelist()) == len(set(previous.namelist()))
                and set(names) == set(previous.namelist()),
                "Office parts differ from the rendered source")
        require(all(actual.read(name) == previous.read(name) for name in names),
                "Office content differs from the rendered source")
    return {"scope": "all_uncompressed_office_parts_identical",
            "parts": len(names), "source_sha256": sha256(source),
            "rendered_source_sha256": sha256(rendered_source), "passed": True}


def check_pair(directory, *, check_tool_hashes=True):
    directory = Path(directory).resolve()
    evidence, provenance = load_evidence()
    topics = build_topics(evidence)
    manifest = load_json(directory / "materials_manifest.json")
    require(manifest["display_status"] == STATUS
            and manifest["scope"] == "detailed_anonymous_office_materials_not_competition_acceptance",
            "material status or scope changed")
    require(all(manifest.get(name) is False for name in (
        "native_latency_passed", "current_v11_verified", "independent_blind",
        "competition_ready", "pdf_delivered", "packaged", "uploaded",
    )), "material manifest makes an unsupported acceptance or delivery claim")
    require(set(manifest["output_raw_sha256"]) == {
        DOCX_NAME, PPTX_NAME, "evidence_provenance.json", "content_index.json",
    }, "material output inventory changed")
    if manifest["word"].get("page_map_from_render"):
        validate_word_layout({
            "scope": "word_directory_fallbacks_from_measured_render",
            "layout_source_lf_sha256": {
                name: manifest["document_tool_lf_sha256"].get(
                    (Path(__file__).parent / name).relative_to(ROOT).as_posix())
                for name in layout_source_hashes()
            },
            "page_count": manifest["word"]["expected_pages"],
            "section_pages": manifest["word"]["section_pages"],
        }, topics)
    require(load_json(directory / "evidence_provenance.json") == provenance,
            "document provenance changed or no longer matches measured runtime")
    for name, expected in manifest["output_raw_sha256"].items():
        path = (directory / name).resolve()
        require(path.is_relative_to(directory) and sha256(path) == expected,
                "material output changed: " + name)
    if check_tool_hashes:
        tools = {path.relative_to(ROOT).as_posix(): lf_sha256(path)
                 for path in sorted(Path(__file__).parent.glob("*.py"))}
        require(tools == manifest["document_tool_lf_sha256"], "document tool snapshot changed")
    require(not any(path.suffix.lower() in (".pdf", ".zip") for path in directory.rglob("*")),
            "this delivery must not contain PDF or ZIP")
    require(sorted(path.name for path in directory.glob("*.docx")) == [DOCX_NAME]
            and sorted(path.name for path in directory.glob("*.pptx")) == [PPTX_NAME],
            "expected exactly one Word and one project report")
    for name in (DOCX_NAME, PPTX_NAME):
        result = inspect_ooxml(directory / name, expected_slide_count=manifest["presentation"]["slides"])
        require(result["passed"], "invalid Office structure: " + repr(result["issues"]))
    word_text = package_text(directory / DOCX_NAME)
    ppt_text = package_text(directory / PPTX_NAME)
    for text in (word_text, ppt_text):
        for required in (STATUS, "待复测风险项", "HashEmbedding", "48.145", "13,390.615",
                         "MemGPT", "Mem0", "MemoryOS", "LongMemEval", "nearest-rank",
                         "46.15%", "31.71%", "2026年9月15日24:00", "5–10", "200M"):
            require(normalized(required) in text, "missing disclosure or official rule: " + required)
    for t in topics:
        require(normalized(t["title"]) in word_text, "missing Word section: " + t["number"])
        for _, paragraph in t["paragraphs"]:
            require(normalized(paragraph) in word_text, "missing Word detail: " + t["number"])
        if t["table"]:
            for row in (t["table"][0], *t["table"][1]):
                for value in row:
                    require(normalized(value) in word_text, "missing Word table content: " + t["number"])
        if t["slide"]:
            for _, paragraph in t["paragraphs"]:
                require(normalized(paragraph) in ppt_text, "missing presenter notes: " + t["number"])
    doc = Document(directory / DOCX_NAME)
    require(len(doc.inline_shapes) == len(DIAGRAMS) == manifest["word"]["figures"],
            "Word diagram count changed")
    prs = Presentation(directory / PPTX_NAME)
    require(len(prs.slides) == len(presentation_topics(topics)) + 2 == manifest["presentation"]["slides"],
            "slide count changed")
    diagrams = set()
    tables = 0
    for index, slide in enumerate(prs.slides):
        require(bool(slide.notes_slide.notes_text_frame.text.strip()), "missing speaker notes")
        for shape in slide.shapes:
            require(shape.left >= 0 and shape.top >= 0 and shape.width >= 0 and shape.height >= 0
                    and shape.left + shape.width <= prs.slide_width + 8
                    and shape.top + shape.height <= prs.slide_height + 8,
                    f"shape outside slide {index + 1}: {shape.name}")
            require(shape.shape_type != 13, "PPT diagrams must be native editable shapes, not bitmaps")
            if shape.name.startswith("diagram:"):
                diagrams.add(shape.name.split(":")[1])
            tables += bool(shape.has_table)
    require(diagrams == DIAGRAMS.keys() and tables == manifest["presentation"]["native_tables"],
            "editable diagram/table coverage changed")
    return {
        "scope": "source_bound_ooxml_content_anonymity_and_editability",
        "passed": True, "runtime_file_count": provenance["runtime_file_count"],
        "word": manifest["word"], "presentation": manifest["presentation"],
        "office_sha256": {name: sha256(directory / name) for name in (DOCX_NAME, PPTX_NAME)},
        "checks": {
            "source_and_evidence_revalidated": True,
            "output_hashes_match": True,
            "tool_hashes_checked": check_tool_hashes,
            "all_word_sections_paragraphs_and_tables_present": True,
            "presenter_notes_complete": True,
            "identity_and_obsolete_status_scan_passed": True,
            "metadata_anonymous_and_no_external_relationships": True,
            "native_ppt_diagrams_and_tables_verified": True,
            "shapes_inside_canvas": True,
            "no_delivered_pdf_or_zip": True,
        },
        "visual_review_required": True,
        "competition_ready": False,
    }


def contacts(images, output, prefix, *, columns, rows, thumb_size):
    sheets = []
    face = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 16)
    tw, th = thumb_size
    cell_w, cell_h = tw + 16, th + 35
    for offset in range(0, len(images), columns * rows):
        batch = images[offset:offset + columns * rows]
        sheet = Image.new("RGB", (columns * cell_w, math_ceil(len(batch), columns) * cell_h), "#e1e7e8")
        draw = ImageDraw.Draw(sheet)
        for i, image_path in enumerate(batch):
            image = Image.open(image_path).convert("RGB")
            image.thumbnail((tw, th), Image.Resampling.LANCZOS)
            x, y = (i % columns) * cell_w + 8, (i // columns) * cell_h + 8
            sheet.paste(image, (x + (tw - image.width) // 2, y))
            draw.text((x + 4, y + th + 6), f"{prefix} {offset + i + 1:02d}", font=face, fill="#172629")
        target = output / f"{prefix}-contact-{len(sheets) + 1:02d}.jpg"
        sheet.save(target, quality=91)
        sheets.append(target.name)
    return sheets


def math_ceil(number, denominator):
    return (number + denominator - 1) // denominator


def rendered_checks(materials, word_pdf, ppt_dir, output, *, check_tool_hashes=True,
                    ppt_render_source=None):
    import fitz

    materials, word_pdf, ppt_dir, output = map(Path, (materials, word_pdf, ppt_dir, output))
    output = output.resolve()
    private_root = Path.home() / "TiaobeiLab"
    require(output.is_relative_to(private_root.resolve()), "QA rasters must remain in the private lab")
    require(not output.exists(), "refusing to replace an existing render review")
    structural = check_pair(materials, check_tool_hashes=check_tool_hashes)
    equivalent = (verify_render_equivalent(materials / PPTX_NAME, ppt_render_source)
                  if ppt_render_source else None)
    evidence, _ = load_evidence()
    topics = build_topics(evidence)
    output.mkdir(parents=True)
    pdf = fitz.open(word_pdf)
    problems, word_pages, word_images = [], [], []
    expected = structural["word"]["expected_pages"]
    if expected is not None and len(pdf) != expected:
        problems.append(f"word_page_count:{len(pdf)};expected:{expected}")
    section_pages, page_texts = {}, []
    for index, page in enumerate(pdf):
        body_lines = []
        for block in page.get_text("dict")["blocks"]:
            if block["type"] == 0:
                body_lines.extend("".join(span["text"] for span in line["spans"])
                                  for line in block["lines"]
                                  if line["bbox"][1] > 40 and line["bbox"][3] < page.rect.height - 35)
        text = normalized("".join(body_lines))
        page_texts.append(text)
        issues = []
        if index >= 3:
            for t in topics:
                if normalized(t["number"] + t["title"]) in text:
                    if t["number"] in section_pages:
                        issues.append("duplicate_rendered_section:" + t["number"])
                    section_pages[t["number"]] = index + 1
        if len(text) < 100:
            issues.append("near_empty_page")
        for block in page.get_text("dict")["blocks"]:
            if block["type"] != 0:
                continue
            for line in block["lines"]:
                x0, y0, x1, y1 = line["bbox"]
                if x0 < 22 or y0 < 12 or x1 > page.rect.width - 22 or y1 > page.rect.height - 12:
                    issues.append("text_near_or_outside_page_edge")
        image_path = output / f"word-page-{index + 1:02d}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(image_path)
        word_images.append(image_path)
        word_pages.append({"page": index + 1, "characters": len(text), "issues": sorted(set(issues))})
        problems.extend(f"word:{index + 1}:{issue}" for issue in sorted(set(issues)))
    all_body = "".join(page_texts)
    for t in topics:
        if t["number"] not in section_pages:
            problems.append("missing_rendered_section:" + t["number"])
        for label, body in t["paragraphs"]:
            if normalized(body) not in all_body:
                problems.append("missing_rendered_paragraph:" + t["number"] + ":" + label)
        if t["table"]:
            for row in (t["table"][0], *t["table"][1]):
                for value in row:
                    if normalized(value) not in all_body:
                        problems.append("missing_rendered_cell:" + t["number"] + ":" + str(value)[:65])
    if structural["word"].get("page_map_from_render"):
        if section_pages != structural["word"]["section_pages"]:
            problems.append("rendered_section_pages_differ_from_directory_fields")
    write_json(output / "word_layout.json", {
        "scope": "word_directory_fallbacks_from_measured_render",
        "layout_source_lf_sha256": layout_source_hashes(),
        "page_count": len(pdf), "section_pages": section_pages,
    })
    ppt_images = sorted(ppt_dir.glob("*.png"), key=lambda path: int(re.search(r"(\d+)(?=\.png$)", path.name, re.I)[1]))
    if len(ppt_images) != structural["presentation"]["slides"]:
        problems.append("ppt_rendered_slide_count")
    ppt_slides = []
    for index, path in enumerate(ppt_images):
        image = Image.open(path).convert("RGB")
        thumbnail = image.resize((320, 180), Image.Resampling.BILINEAR)
        ink = sum(min(thumbnail.getpixel((x, y))) < 210
                  for y in range(180) for x in range(320)) / (320 * 180)
        stddev = max(ImageStat.Stat(thumbnail).stddev)
        valid = image.size == (1600, 900) and ink >= .009 and stddev > 12
        if not valid:
            problems.append(f"ppt:{index + 1}:blank_or_wrong_size")
        ppt_slides.append({"slide": index + 1, "size": list(image.size),
                           "dark_pixel_fraction": round(ink, 4), "nonblank": valid})
    sheets = contacts(word_images, output, "word", columns=3, rows=4, thumb_size=(260, 368))
    sheets += contacts(ppt_images, output, "ppt", columns=3, rows=3, thumb_size=(360, 203))
    result = {
        "scope": "office_render_text_coverage_layout_and_pixel_checks",
        "structural": structural, "word_rendered_pages": len(pdf),
        "ppt_rendered_slides": len(ppt_images), "word_pages": word_pages, "ppt_slides": ppt_slides,
        "contact_sheets": sheets, "issues": problems, "automated_render_checks_passed": not problems,
        "human_visual_review_recorded": False,
        "ppt_render_source_equivalence": equivalent,
        "private_pdf_is_not_a_deliverable": True, "competition_ready": False,
    }
    write_json(output / "render_checks.json", result)
    print(json.dumps({"pages": len(pdf), "slides": len(ppt_images), "issues": problems,
                      "contact_sheets": sheets}, ensure_ascii=False, indent=2))
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--materials", type=Path, required=True)
    parser.add_argument("--word-render", type=Path)
    parser.add_argument("--ppt-render", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--ppt-render-source", type=Path,
                        help="original PPTX when reusing a render with identical uncompressed parts")
    parser.add_argument("--draft", action="store_true",
                        help="allow a draft made before QA-tool edits; runtime/output hashes remain mandatory")
    args = parser.parse_args(argv)
    if any((args.word_render, args.ppt_render, args.output)):
        require(all((args.word_render, args.ppt_render, args.output)), "provide all three rendering arguments")
        result = rendered_checks(args.materials, args.word_render, args.ppt_render, args.output,
                                 check_tool_hashes=not args.draft,
                                 ppt_render_source=args.ppt_render_source)
        return 0 if result["automated_render_checks_passed"] else 2
    print(json.dumps(check_pair(args.materials, check_tool_hashes=not args.draft), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
