"""Build the detailed anonymous DOCX/PPTX pair from validated frozen evidence."""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timezone
from functools import lru_cache
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as wqn
from docx.shared import Inches as WI, Pt as WP, RGBColor as WC
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement as PX
from pptx.util import Inches as PI, Pt as PP

from docs.materials_20260911.content import CHAPTERS, DATE, STATUS, TITLE, build_topics
from docs.materials_20260911.diagrams import DIAGRAMS, HEIGHT, PALETTE as C, WIDTH, edge_points
from docs.materials_20260911.evidence import ROOT, lf_sha256, load_evidence, load_json, require, sha256

DOCX_NAME = "OS_Agent_技术方案_20260911.docx"
PPTX_NAME = "OS_Agent_项目报告_20260911.pptx"
DEFAULT_OUTPUT = Path("deliverables/submission-materials-20260911")
FONTS = Path("C:/Windows/Fonts")
LATIN = "Arial"
SANS = "Microsoft YaHei"
SERIF = "SimSun"
MONO = "Consolas"
BODY_WIDTH = 6.77
SLIDE_WIDTH, SLIDE_HEIGHT = 13.333333, 7.5
GENERATED_AT = datetime(2026, 9, 11, tzinfo=timezone.utc)


@lru_cache(maxsize=256)
def font(size, bold=False, mono=False):
    filename = "consola.ttf" if mono else "msyhbd.ttc" if bold else "msyh.ttc"
    path = FONTS / filename
    require(path.is_file(), f"missing document font: {filename}")
    return ImageFont.truetype(str(path), max(1, round(size * 4)))


def width(text, size, bold=False, mono=False):
    return font(size, bold, mono).getlength(text) / 4


def wrap(text, max_width, size, bold=False, mono=False):
    lines = []
    for original in str(text).split("\n"):
        current = ""
        tokens = re.findall(r"[A-Za-z0-9_]+(?:[./:%+\-][A-Za-z0-9_]+)*|[^\x00-\x7f]|\s+|.", original)
        for token in tokens:
            if width(current + token, size, bold, mono) <= max_width:
                current += token
                continue
            if current.strip():
                lines.append(current.rstrip())
                current = ""
            if width(token, size, bold, mono) <= max_width:
                current = token.lstrip()
            else:
                for character in token:
                    if current and width(current + character, size, bold, mono) > max_width:
                        lines.append(current)
                        current = ""
                    current += character
        lines.append(current.rstrip())
    return lines


def fitted(text, max_width, max_height, size, minimum, *, bold=False, mono=False, spacing=1.18):
    while size >= minimum:
        lines = wrap(text, max_width, size, bold, mono)
        if len(lines) * size * spacing <= max_height:
            return size, lines
        size -= .5
    raise ValueError(f"text will not fit safely: {text[:90]}")


def diagram_fonts(node):
    if node.get("lane"):
        return 32, 29
    title_size = 32
    while width(node["title"], title_size, True) > node["w"] - 80:
        title_size -= .5
    body_size = 29
    while max(width(line, body_size) for line in node["body"].splitlines()) > node["w"] - 36:
        body_size -= .5
    require(title_size >= 23 and body_size >= 23, "diagram label needs a shorter phrase")
    return title_size, body_size


def diagram_png(key):
    spec = DIAGRAMS[key]
    scale = 2
    image = Image.new("RGB", (WIDTH * scale, HEIGHT * scale), "white")
    draw = ImageDraw.Draw(image)

    def line(a, b, fill, stroke=2):
        draw.line([tuple(v * scale for v in a), tuple(v * scale for v in b)],
                  fill=fill, width=round(stroke * scale))

    def text(value, x, y, size, fill, bold=False):
        # Fonts use four times the point/design unit; reduce to the raster scale.
        face = ImageFont.truetype(str(FONTS / ("msyhbd.ttc" if bold else "msyh.ttc")), round(size * scale))
        draw.text((x * scale, y * scale), value, font=face, fill=fill, anchor="lt")

    for start, end in spec["edges"]:
        a, b = edge_points(spec["nodes"][start], spec["nodes"][end])
        line(a, b, "#" + C["muted"], 2.6)
        angle = math.atan2(b[1] - a[1], b[0] - a[0])
        points = [b, (b[0] - 12 * math.cos(angle - .48), b[1] - 12 * math.sin(angle - .48)),
                  (b[0] - 12 * math.cos(angle + .48), b[1] - 12 * math.sin(angle + .48))]
        draw.polygon([(x * scale, y * scale) for x, y in points], fill="#" + C["muted"])
    for node in spec["nodes"]:
        x, y, w, h = (node[key] for key in ("x", "y", "w", "h"))
        tone = node["tone"]
        pale = "pale" if tone == "teal" else tone + "_pale"
        draw.rectangle((x * scale, y * scale, (x + w) * scale, (y + h) * scale),
                       fill="#" + C[pale], outline="#" + C[tone], width=2 * scale)
        ts, bs = diagram_fonts(node)
        if node.get("lane"):
            text(node["title"], x + 22, y + 23, ts, "#" + C[tone], True)
            text(node["body"], x + 235, y + 24, bs, "#" + C["ink"])
        else:
            text(node["title"], x + 18, y + 18, ts, "#" + C[tone], True)
            text(node["step"], x + w - 49, y + 20, 23, "#" + C["muted"])
            for i, value in enumerate(node["body"].splitlines()):
                text(value, x + (w - width(value, bs)) / 2, y + 78 + i * 38, bs, "#" + C["ink"])
    output = BytesIO()
    image.save(output, format="PNG", dpi=(360, 360), optimize=True)
    output.seek(0)
    return output


def word_font(run, size=None, bold=False, color=None, face=None):
    run.font.name = face or LATIN
    if size is not None:
        run.font.size = WP(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = WC.from_string(C.get(color, color))
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    fonts.set(wqn("w:eastAsia"), SANS if bold or face == SANS else SERIF)
    fonts.set(wqn("w:ascii"), face or LATIN)
    fonts.set(wqn("w:hAnsi"), face or LATIN)
    return run


def shade(element, fill):
    node = OxmlElement("w:shd")
    node.set(wqn("w:fill"), C.get(fill, fill))
    element.append(node)


def bookmark(paragraph, name, identifier):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(wqn("w:id"), str(identifier))
    start.set(wqn("w:name"), name)
    end.set(wqn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def field(paragraph, instruction, fallback, size=9):
    simple = OxmlElement("w:fldSimple")
    simple.set(wqn("w:instr"), instruction)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(wqn("w:val"), str(round(size * 2)))
    props.append(sz)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = str(fallback)
    run.append(value)
    simple.append(run)
    paragraph._p.append(simple)


def word_paragraph(doc, text="", *, size=11, bold=False, color="ink", after=6, before=0,
                   face=None, align=None, line=1.2, keep=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = WP(before), WP(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    fmt.widow_control = True
    if align is not None:
        p.alignment = align
    word_font(p.add_run(text), size, bold, color, face)
    return p


def table_ratios(number, count):
    special = {
        "2.4": (0.20, .44, .36), "3.1": (.25, .43, .32),
        "11.1": (.23, .32, .45), "11.2": (.23, .39, .38),
        "12.2": (.18, .42, .40), "12.3": (.28, .13, .34, .25),
        "12.4": (.20, .43, .37), "13.1": (.32, .25, .20, .23),
        "13.2": (.28, .18, .18, .18, .18),
        "13.3": (.25, .15, .13, .13, .13, .21),
        "13.4": (.18, .18, .32, .32),
        "13.5": (.29, .17, .17, .37), "14.2": (.34, .40, .26),
        "16.1": (.26, .39, .35), "18.2": (.16, .52, .32),
        "18.3": (.30, .70), "18.4": (.09, .63, .28), "18.5": (.08, .50, .42),
    }
    if number in special:
        return special[number]
    if count == 3:
        return .24, .37, .39
    return (1 / count,) * count


def word_table(doc, table_spec, number, *, size=9.6):
    headers, rows = table_spec
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    ratios = table_ratios(number, len(headers))
    for i, ratio in enumerate(ratios):
        table.columns[i].width = WI(BODY_WIDTH * ratio)
    for row_index, values in enumerate([headers, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        props.append(no_split)
        if row_index == 0:
            props.append(OxmlElement("w:tblHeader"))
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = WI(BODY_WIDTH * ratios[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            shade(tc, "teal" if row_index == 0 else "pale" if row_index % 2 else "paper")
            margins = OxmlElement("w:tcMar")
            for side, amount in (("top", 70), ("bottom", 70), ("left", 90), ("right", 90)):
                part = OxmlElement("w:" + side)
                part.set(wqn("w:w"), str(amount))
                part.set(wqn("w:type"), "dxa")
                margins.append(part)
            tc.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = WP(0)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = True
            if number == "18.3" and row_index and column == 1:
                value = str(value)[:32] + "\n" + str(value)[32:]
            word_font(p.add_run(str(value)), size, row_index == 0,
                      "paper" if row_index == 0 else "ink")
    return table


def setup_word(page_count=1):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = WI(8.27), WI(11.69)
    section.top_margin = section.bottom_margin = WI(.65)
    section.left_margin = section.right_margin = WI(.75)
    section.header_distance = section.footer_distance = WI(.28)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = WP(12)
    normal.element.get_or_add_rPr().rFonts.set(wqn("w:eastAsia"), SERIF)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = WP(6)
    for name, size in (("Heading 1", 19), ("Heading 2", 15)):
        style = doc.styles[name]
        style.font.name = SANS
        style.font.size = WP(size)
        style.font.bold = True
        style.font.color.rgb = WC.from_string(C["teal"] if name == "Heading 1" else C["ink"])
        style.element.get_or_add_rPr().rFonts.set(wqn("w:eastAsia"), SANS)
        style.paragraph_format.space_before = WP(0)
        style.paragraph_format.space_after = WP(9 if name == "Heading 2" else 14)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = WP(0)
    word_font(header.add_run("OS Agent 多源融合记忆优化系统  |  技术方案"), 8, color="muted")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    word_font(footer.add_run("已完成初版，待赛事验收   |   "), 8, color="muted")
    field(footer, "PAGE", "1", 8)
    word_font(footer.add_run(" / "), 8, color="muted")
    field(footer, "NUMPAGES", str(page_count), 8)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(wqn("w:val"), "true")
    doc.settings.element.append(update_fields)
    doc.core_properties.title = TITLE + " 技术方案"
    doc.core_properties.subject = "技术文档、用户手册与效果验证报告"
    doc.core_properties.author = doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "OS Agent; memory; anonymous; public regression"
    doc.core_properties.comments = ""
    doc.core_properties.created = doc.core_properties.modified = GENERATED_AT
    return doc


def word_cover(doc, e):
    word_paragraph(doc, "挑战杯 · 揭榜挂帅  |  麒麟赛题", size=12, bold=True, color="teal",
                   before=35, after=32, align=WD_ALIGN_PARAGRAPH.CENTER)
    word_paragraph(doc, "OS Agent\n多源融合记忆优化系统", size=27, bold=True,
                   after=25, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.3)
    word_paragraph(doc, "技术方案", size=23, bold=True, after=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    word_paragraph(doc, "含用户手册与效果验证报告", size=14, color="muted",
                   after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
    word_paragraph(doc, STATUS, size=13, bold=True, color="teal",
                   after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    word_paragraph(doc, "材料版本 2026-09-11  |  匿名评审稿", size=10, color="muted",
                   after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    word_table(doc, (["本版范围", "说明"], [
        ["完整技术说明", "18章、67节；方案、算法、实现、部署、操作、评测与证据索引"],
        ["当前实现", "focused_evidence_v2；六类事件、偏好版本、记忆流转、冲突与遗忘"],
        ["量化证据", f"当前Windows公开合成回归；{e['tests']}项测试无失败、{e['skipped']}项跳过"],
        ["证据边界", "HashEmbedding不代表真实SDK；非独立盲测或赛事认证"],
        ["待复测风险", "真实SDK时延、当前源码V11复验及真实Agent/Judge与授权案例"],
    ]), "cover", size=10.5)
    word_paragraph(doc, "本次提供可编辑 Word；不制作交付 PDF，不打包、不上传。正式外层命名与提交由负责人审阅后办理。",
                   size=10, color="muted", before=15, align=WD_ALIGN_PARAGRAPH.CENTER)


def word_navigation(doc, topics, page_map):
    doc.add_page_break()
    word_paragraph(doc, "目录与阅读路径", size=22, bold=True, after=14)
    word_paragraph(doc, "第2–10章为方案与实现，第11章为相关工作和创新，第12–14章为效果验证报告，第15–18章为用户手册、接口与验收资料。",
                   size=11, after=17)
    first = {chapter: next((i, t) for i, t in enumerate(topics) if t["number"].split(".")[0] == chapter)
             for chapter in CHAPTERS}
    for chapter, title in CHAPTERS.items():
        index, target = first[chapter]
        p = word_paragraph(doc, "", size=11, after=7)
        p.paragraph_format.tab_stops.add_tab_stop(WI(BODY_WIDTH - .35))
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(wqn("w:anchor"), "sec_" + target["number"].replace(".", "_"))
        run = OxmlElement("w:r")
        props, color = OxmlElement("w:rPr"), OxmlElement("w:color")
        color.set(wqn("w:val"), C["teal"])
        props.append(color)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = f"第{chapter}章  {title}"
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)
        p.add_run("\t")
        field(p, "PAGEREF sec_" + target["number"].replace(".", "_") + " \\h",
              page_map.get(target["number"], index + 4), 11)
    word_paragraph(doc, "目录页码按本版排版设置；在Word中修改内容后，可更新域以刷新页码。证据编号W/R/S/C/E/N/D/H/SDK见18.2节。",
                   size=9.5, color="muted", before=15)
    doc.add_page_break()
    word_paragraph(doc, "方案框图索引", size=22, bold=True, after=12)
    word_paragraph(doc, "框图按实际模块和授权边界绘制。技术方案内为高清图片；项目报告使用可编辑的原生形状、文字和连接线。",
                   size=11, after=16)
    count = 0
    for index, t in enumerate(topics):
        if not t["diagram"]:
            continue
        count += 1
        p = word_paragraph(doc, f"图{count:02d}  {DIAGRAMS[t['diagram']]['title']}", size=10, after=4)
        p.paragraph_format.tab_stops.add_tab_stop(WI(BODY_WIDTH - .35))
        p.add_run("\t")
        field(p, "PAGEREF sec_" + t["number"].replace(".", "_") + " \\h",
              page_map.get(t["number"], index + 4), 10)
    word_paragraph(doc, "箭头表示对应图中的处理顺序或依赖，不代表未实现的自动同步；条件失败可提前退出。模型、授权与正式验收边界在相应章节单列。",
                   size=9.5, color="muted", before=15)


def make_word(path, topics, e, layout=None):
    page_map = layout["section_pages"] if layout else {}
    doc = setup_word(layout["page_count"] if layout else 1)
    word_cover(doc, e)
    word_navigation(doc, topics, page_map)
    figures, seen_chapters = 0, set()
    for index, t in enumerate(topics):
        chapter = t["number"].split(".")[0]
        if chapter not in seen_chapters:
            doc.add_page_break()
            kicker = doc.add_paragraph(f"第{chapter}章  {CHAPTERS[chapter]}")
            kicker.style = doc.styles["Heading 1"]
            seen_chapters.add(chapter)
        heading = doc.add_paragraph(f"{t['number']}  {t['title']}", "Heading 2")
        if not t["number"].endswith(".1"):
            heading.paragraph_format.space_before = WP(17)
        bookmark(heading, "sec_" + t["number"].replace(".", "_"), index + 1)
        word_paragraph(doc, t["lead"], size=11.5, color="teal", after=11, line=1.22, keep=True)
        if t["diagram"]:
            figures += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = WP(4)
            p.paragraph_format.keep_with_next = True
            picture = p.add_run().add_picture(diagram_png(t["diagram"]), width=WI(BODY_WIDTH))
            spec = DIAGRAMS[t["diagram"]]
            picture._inline.docPr.set("descr", spec["title"] + "；" + " → ".join(
                node["title"] + "：" + node["body"].replace("\n", "，") for node in spec["nodes"]))
            word_paragraph(doc, f"图{figures:02d}  {spec['title']}。{spec['note']}",
                           size=9, color="muted", after=10, line=1.15)
        for label, body in t["paragraphs"]:
            p = word_paragraph(doc, "", size=12, after=6, line=1.20,
                               align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            word_font(p.add_run(label + "。"), 12, True)
            word_font(p.add_run(body), 12)
        if t["table"]:
            word_table(doc, t["table"], t["number"], size=9.2 if len(t["table"][1]) > 8 else 9.6)
        if t["code"]:
            p = word_paragraph(doc, t["code"], size=9.2, face=MONO, line=1.18, after=8)
            shade(p._p.get_or_add_pPr(), "blue_pale")
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = WI(.08)
            p.paragraph_format.keep_together = True
        if t["warning"]:
            word_paragraph(doc, t["warning"], size=9.5, bold=True, color="amber", after=8)
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        source = word_paragraph(doc, "依据：" + t["source"], size=8.2, color="muted", before=3, after=0)
        source.paragraph_format.keep_together = True
    doc.save(path)
    return {"expected_pages": layout["page_count"] if layout else None,
            "section_pages": page_map, "page_map_from_render": bool(layout),
            "topics": len(topics), "chapters": len(CHAPTERS),
            "figures": figures, "native_tables": len(doc.tables)}


def ppt_font(run, size, bold=False, color="ink", face=SANS):
    run.font.name = face
    run.font.size = PP(size)
    run.font.bold = bold
    run.font.color.rgb = PC.from_string(C.get(color, color))
    props = run._r.get_or_add_rPr()
    props.set("lang", "zh-CN")
    for tag in ("a:ea", "a:cs"):
        old = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag[2:])
        if old is not None:
            props.remove(old)
        node = PX(tag)
        node.set("typeface", face if face != MONO else SANS)
        props.append(node)


def textbox(slide, text, x, y, w, h, size, *, bold=False, color="ink", align=PP_ALIGN.LEFT,
            vertical=MSO_ANCHOR.TOP, face=SANS, minimum=None, spacing=1.15, name=None):
    require(x >= 0 and y >= 0 and x + w <= SLIDE_WIDTH + .001
            and y + h <= SLIDE_HEIGHT + .001, "textbox outside slide")
    size, lines = fitted(text, w * 72, h * 72, size, size if minimum is None else minimum,
                         bold=bold, mono=face == MONO, spacing=spacing)
    shape = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = vertical
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = PP(size * spacing)
        p.space_before = p.space_after = PP(0)
        ppt_font(p.add_run(), size, bold, color, face)
        p.runs[-1].text = line
    return shape


def rect(slide, x, y, w, h, fill, stroke=None, stroke_width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PI(x), PI(y), PI(w), PI(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PC.from_string(C.get(fill, fill))
    if stroke is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = PC.from_string(C.get(stroke, stroke))
        shape.line.width = PP(stroke_width)
    return shape


def connector(slide, a, b, stroke="muted", arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PI(a[0]), PI(a[1]), PI(b[0]), PI(b[1]))
    shape.line.color.rgb = PC.from_string(C[stroke])
    shape.line.width = PP(1.4)
    if arrow:
        end = PX("a:tailEnd")
        end.set("type", "triangle")
        shape.line._get_or_add_ln().append(end)
    return shape


def ppt_diagram(slide, key, x=.992, y=1.80, w=11.35):
    spec = DIAGRAMS[key]
    scale = w / WIDTH
    for start, end in spec["edges"]:
        a, b = edge_points(spec["nodes"][start], spec["nodes"][end])
        connector(slide, (x + a[0] * scale, y + a[1] * scale),
                  (x + b[0] * scale, y + b[1] * scale))
    for index, node in enumerate(spec["nodes"]):
        nx, ny, nw, nh = (node[field] for field in ("x", "y", "w", "h"))
        tone = node["tone"]
        pale = "pale" if tone == "teal" else tone + "_pale"
        box = rect(slide, x + nx * scale, y + ny * scale, nw * scale, nh * scale, pale, tone, 1.0)
        box.name = f"diagram:{key}:node:{index + 1}"
        ts, bs = diagram_fonts(node)
        if node.get("lane"):
            textbox(slide, node["title"], x + (nx + 22) * scale, y + (ny + 20) * scale,
                    205 * scale, 48 * scale, ts * scale * 72, bold=True, color=tone)
            textbox(slide, node["body"], x + (nx + 235) * scale, y + (ny + 21) * scale,
                    901 * scale, 49 * scale, bs * scale * 72, minimum=25 * scale * 72)
        else:
            textbox(slide, node["title"], x + (nx + 18) * scale, y + (ny + 18) * scale,
                    (nw - 80) * scale, 46 * scale, ts * scale * 72, bold=True, color=tone)
            textbox(slide, node["step"], x + (nx + nw - 49) * scale, y + (ny + 20) * scale,
                    40 * scale, 38 * scale, 23 * scale * 72, color="muted")
            textbox(slide, node["body"], x + (nx + 18) * scale, y + (ny + 76) * scale,
                    (nw - 36) * scale, 78 * scale, bs * scale * 72,
                    align=PP_ALIGN.CENTER, spacing=1.20)
    return y + HEIGHT * scale


def notes(slide, t, extra=""):
    lines = [t["number"] + " " + t["title"], t["lead"], ""]
    lines.extend(label + "：" + body for label, body in t["paragraphs"])
    if t["diagram"]:
        spec = DIAGRAMS[t["diagram"]]
        lines.extend(["", "框图：" + spec["title"], spec["note"]])
        lines.extend(node["step"] + " " + node["title"] + "：" + node["body"].replace("\n", "；")
                     for node in spec["nodes"])
    if t["table"]:
        lines.extend(["", " | ".join(t["table"][0])])
        lines.extend(" | ".join(str(value) for value in row) for row in t["table"][1])
    if t["code"]:
        lines.extend(["", t["code"]])
    lines.extend(["", t["warning"], "依据：" + t["source"], extra])
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def points_row(slide, points, y=6.55):
    for i, value in enumerate(points[:3]):
        x = .67 + i * 4.03
        rect(slide, x, y + .03, .035, .48, "teal")
        textbox(slide, value, x + .13, y, 3.78, .57, 15.5, minimum=13.5, spacing=1.16)


def slide_base(prs, number, total, title, chapter, lead, source):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PC.from_string(C["paper"])
    rect(slide, .62, .30, .07, .30, "teal")
    textbox(slide, chapter, .84, .29, 11.5, .24, 11, color="teal", bold=True)
    textbox(slide, title, .64, .67, 12.00, .52, 29, bold=True, minimum=26)
    textbox(slide, lead, .67, 1.29, 12.00, .49, 17, minimum=15.5, color="muted")
    connector(slide, (.66, 7.21), (12.65, 7.21), "line", False)
    textbox(slide, "依据：" + source, .67, 7.27, 11.38, .20, 8.5, minimum=7.5, color="muted")
    textbox(slide, f"{number:02d} / {total:02d}", 12.06, 7.27, .61, .20, 9,
            align=PP_ALIGN.RIGHT, color="muted")
    return slide


def ppt_table(slide, spec, number, x=.67, y=2.02, w=12.00, max_h=4.22):
    headers, rows = spec
    ratios = table_ratios(number, len(headers))
    size = 17 if len(headers) < 5 else 16
    if number == "13.4":
        size = 16.5
    while size >= 14:
        wrapped = []
        heights = []
        for i, values in enumerate([headers, *rows]):
            cells = [wrap(str(value), w * ratios[j] * 72 - 16, size, i == 0)
                     for j, value in enumerate(values)]
            wrapped.append(cells)
            heights.append(max(len(cell) for cell in cells) * size * 1.15 / 72 + .17)
        if sum(heights) <= max_h:
            break
        size -= .5
    require(sum(heights) <= max_h and size >= 14, f"table {number} requires a separate slide")
    graphic = slide.shapes.add_table(len(rows) + 1, len(headers), PI(x), PI(y), PI(w), PI(sum(heights)))
    graphic.name = "editable_table:" + number
    table = graphic.table
    for column, ratio in enumerate(ratios):
        table.columns[column].width = PI(w * ratio)
    for i, values in enumerate(wrapped):
        table.rows[i].height = PI(heights[i])
        for j, lines in enumerate(values):
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = PI(.10)
            cell.margin_top = cell.margin_bottom = PI(.065)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PC.from_string(C["teal" if i == 0 else "pale" if i % 2 else "paper"])
            cell.text_frame.clear()
            cell.text_frame.word_wrap = False
            for k, value in enumerate(lines):
                p = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                p.line_spacing = PP(size * 1.15)
                p.space_before = p.space_after = PP(0)
                run = p.add_run()
                run.text = value
                ppt_font(run, size, i == 0, "paper" if i == 0 else "ink")
            props = cell._tc.get_or_add_tcPr()
            for side in ("lnL", "lnR", "lnT", "lnB"):
                border = PX("a:" + side)
                border.set("w", "6350")
                solid, color = PX("a:solidFill"), PX("a:srgbClr")
                color.set("val", "FFFFFF")
                solid.append(color)
                border.append(solid)
                props.append(border)
    return y + sum(heights)


def presentation_topics(topics):
    result = []
    for t in topics:
        if not t["slide"]:
            continue
        if t["number"] == "13.4":
            for index, limits in enumerate(((0, 4), (4, 12))):
                part = deepcopy(t)
                part["title"] += "：原120请求" if index == 0 else "：约束48请求"
                part["table"] = (t["table"][0], t["table"][1][slice(*limits)])
                result.append(part)
        else:
            result.append(t)
    return result


def make_ppt(path, topics, e):
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(SLIDE_WIDTH), PI(SLIDE_HEIGHT)
    prs.core_properties.title = TITLE + " 项目报告"
    prs.core_properties.subject = "方案架构、实现机制与效果验证"
    prs.core_properties.author = prs.core_properties.last_modified_by = ""
    prs.core_properties.keywords = "OS Agent; memory; anonymous; public regression"
    prs.core_properties.comments = ""
    prs.core_properties.created = prs.core_properties.modified = GENERATED_AT
    slide_topics = presentation_topics(topics)
    total = len(slide_topics) + 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, .72, .67, .08, .36, "teal")
    textbox(slide, "挑战杯 · 揭榜挂帅  |  麒麟赛题  |  匿名项目报告", .98, .69, 11.5, .33,
            15, bold=True, color="teal")
    textbox(slide, "OS Agent\n多源融合记忆优化系统", .79, 1.43, 11.75, 1.65,
            42, bold=True, spacing=1.24)
    textbox(slide, "方案架构 · 关键实现 · 效果验证 · 部署与验收", .84, 3.39, 11.7, .42,
            21, color="muted")
    for i, (title, body) in enumerate((
        ("记忆治理", "多源接入、动态偏好\n版本、冲突与可确认遗忘"),
        ("证据检索", "聚焦问题、保护主体\n补齐章节、检查正文支持"),
        ("可核验证据", f"{e['tests']}项测试无失败\n168个不同复杂请求分层回归"),
    )):
        x = .85 + i * 4.07
        rect(slide, x, 4.48, 3.7, .035, "teal" if i != 1 else "blue")
        textbox(slide, title, x, 4.75, 3.7, .36, 20, bold=True, color="teal" if i != 1 else "blue")
        textbox(slide, body, x, 5.27, 3.7, .71, 18)
    textbox(slide, STATUS + "  |  2026-09-11", .86, 6.51, 11.6, .34, 16, bold=True, color="teal")
    textbox(slide, "Windows / HashEmbedding 公开合成回归；真实SDK为待复测风险项，正式赛事验收另行完成。",
            .86, 7.04, 11.55, .26, 11, color="muted")
    slide.notes_slide.notes_text_frame.text = (
        TITLE + "\n" + STATUS + "\n本报告与技术方案使用相同当前源码和固定原始证据。"
        "项目并非独立盲测或原生SDK通过，真实授权案例及目标V11仍须独立留证。"
        "本次只提供PPTX与DOCX，不打包或上传。"
    )
    slide = slide_base(prs, 2, total, "报告结构与重点", "阅读路径",
                       "报告按技术方案章节编号组织；完整参数、用户手册和原始证据索引保留在配套Word。",
                       "官方提交规则；配套技术方案目录")
    sections = (
        ("01", "目标与总体方案", "问题、数据模型、分层架构与两条实现路线", "第1–2章"),
        ("02", "关键机制与适配", "多源、偏好、流转、冲突、检索、端侧与安全", "第3–10章"),
        ("03", "创新与效果验证", "相关工作、固定分母、对照、干扰与全矩阵时延", "第11–13章"),
        ("04", "应用、操作与验收", "合成场景、真实案例边界、用户操作及独立评测", "第14–18章"),
    )
    for i, (index, title, body, chapter) in enumerate(sections):
        y = 2.10 + i * 1.06
        textbox(slide, index, .79, y, .7, .45, 26, bold=True, color="teal")
        textbox(slide, title, 1.67, y, 7.8, .39, 22, bold=True)
        textbox(slide, body, 1.69, y + .48, 9.8, .32, 17, color="muted")
        textbox(slide, chapter, 10.55, y + .02, 1.85, .32, 16, color="blue", align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = (
        "官方要求逐项映射：PPT系统展示；Word包含技术文档、用户手册和效果验证报告。"
        "第11章支撑创新性论证，第12至13章交代性能可信度。评分由赛事评审决定，不自行评分。"
    )
    for number, t in enumerate(slide_topics, start=3):
        chapter = t["number"].split(".")[0]
        slide = slide_base(prs, number, total, f"{t['number']}  {t['title']}",
                           f"第{chapter}章  {CHAPTERS[chapter]}", t["lead"], t["source"])
        if t["diagram"]:
            bottom = ppt_diagram(slide, t["diagram"])
            spec = DIAGRAMS[t["diagram"]]
            textbox(slide, spec["note"], .80, bottom + .03, 11.75, .37,
                    12.5, minimum=11.5, color="amber" if t["warning"] else "muted",
                    align=PP_ALIGN.CENTER)
        elif t["table"]:
            bottom = ppt_table(slide, t["table"], t["number"])
            if t["warning"]:
                textbox(slide, t["warning"], .77, max(bottom + .16, 5.85), 11.80, .45,
                        13, minimum=11.5, color="amber")
        elif t["code"]:
            rect(slide, .67, 2.02, 12.00, 4.10, "blue_pale")
            textbox(slide, t["code"], .88, 2.25, 11.55, 3.70, 19,
                    minimum=17, face=MONO, spacing=1.20)
        else:
            for i, point in enumerate(t["points"]):
                y = 2.10 + i * 1.40
                textbox(slide, f"0{i + 1}", .78, y, .71, .43, 25, bold=True, color="teal")
                textbox(slide, point, 1.66, y, 10.80, .43, 21, minimum=20, bold=True)
                textbox(slide, t["paragraphs"][i][1], 1.68, y + .52, 10.72, .77,
                        16.5, minimum=15.5, color="muted")
        if t["diagram"] or t["table"] or t["code"]:
            points_row(slide, t["points"])
        else:
            textbox(slide, "证据边界：公开合成Windows回归；真实SDK、V11与独立Agent/Judge另验。",
                    .81, 6.65, 11.74, .34, 14.5, color="amber")
        notes(slide, t)
    prs.save(path)
    return {"slides": total, "editable_diagrams": sum(bool(t["diagram"]) for t in slide_topics),
            "native_tables": sum(bool(t["table"]) for t in slide_topics),
            "topic_numbers": [t["number"] for t in slide_topics]}


def write_json(path, data):
    with path.open("x", encoding="utf-8") as stream:
        stream.write(json.dumps(data, ensure_ascii=False, indent=2) + "\n")


def layout_source_hashes():
    return {name: lf_sha256(Path(__file__).parent / name)
            for name in ("content.py", "diagrams.py", "generate.py")}


def validate_word_layout(layout, topics):
    require(layout.get("scope") == "word_directory_fallbacks_from_measured_render"
            and layout.get("layout_source_lf_sha256") == layout_source_hashes()
            and set(layout.get("section_pages", {})) == {t["number"] for t in topics}
            and type(layout.get("page_count")) is int and layout["page_count"] > 3,
            "Word page map does not match the content and layout source")
    pages = [layout["section_pages"][t["number"]] for t in topics]
    require(all(type(page) is int and 4 <= page <= layout["page_count"] for page in pages)
            and pages == sorted(pages),
            "Word page map has invalid or unordered page numbers")


def generate(output, word_layout=None):
    output = Path(output).resolve()
    require(output.is_relative_to(ROOT), "output must remain in this project")
    require(not output.exists(), "refusing to overwrite an existing material directory")
    evidence, provenance = load_evidence()
    topics = build_topics(evidence)
    require(set(t["diagram"] for t in topics if t["diagram"]) == DIAGRAMS.keys(),
            "unreferenced or missing diagram")
    layout = load_json(word_layout) if word_layout else None
    if layout is not None:
        validate_word_layout(layout, topics)
    output.mkdir(parents=True)
    word = make_word(output / DOCX_NAME, topics, evidence, layout)
    ppt = make_ppt(output / PPTX_NAME, topics, evidence)
    write_json(output / "evidence_provenance.json", provenance)
    index = {
        "word": [{"section": t["number"], "title": t["title"],
                  "expected_page": word["section_pages"].get(t["number"]),
                  "diagram": t["diagram"]} for i, t in enumerate(topics)],
        "presentation": [{"slide": i + 3, "section": t["number"], "title": t["title"],
                          "diagram": t["diagram"]} for i, t in enumerate(presentation_topics(topics))],
    }
    write_json(output / "content_index.json", index)
    tools = {path.relative_to(ROOT).as_posix(): lf_sha256(path)
             for path in sorted(Path(__file__).parent.glob("*.py"))}
    manifest = {
        "schema_version": 1, "display_status": STATUS, "date": DATE,
        "scope": "detailed_anonymous_office_materials_not_competition_acceptance",
        "word": word, "presentation": ppt,
        "document_tool_lf_sha256": tools,
        "output_raw_sha256": {path.name: sha256(path) for path in sorted(output.iterdir()) if path.is_file()},
        "source_validation_passed": all(provenance["checks"].values()),
        "visual_render_verified": False,
        "native_latency_passed": False, "current_v11_verified": False,
        "independent_blind": False, "competition_ready": False,
        "pdf_delivered": False, "packaged": False, "uploaded": False,
    }
    write_json(output / "materials_manifest.json", manifest)
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--word-layout", type=Path, help="page references derived from a matching WPS render")
    args = parser.parse_args(argv)
    manifest = generate(args.output_dir, args.word_layout)
    print(json.dumps({key: manifest[key] for key in ("word", "presentation", "source_validation_passed")},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
