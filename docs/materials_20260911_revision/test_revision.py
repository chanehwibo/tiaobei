"""Focused revision checks, separate from runtime tests and performance samples."""
from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx import Document
from pptx import Presentation
from pptx.util import Pt

from docs.materials_20260911 import content, evidence
from docs.materials_20260911.diagrams import DIAGRAMS
from docs.materials_20260911_revision import analysis, generate, qa
from docs.materials_20260911_revision.slide_content import build_slides
from docs.materials_20260911_revision.word_content import CHAPTERS, FOCUS, build_topics


class RevisionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.evidence, cls.provenance = evidence.load_evidence()
        cls.analysis = analysis.derive_analysis(cls.provenance)
        cls.topics = build_topics(cls.evidence, cls.analysis)
        cls.slides = build_slides(cls.evidence, cls.analysis)
        cls.temporary = tempfile.TemporaryDirectory(prefix="materials-revision-tests-", dir=generate.PRIVATE_ROOT)
        cls.addClassCleanup(cls.temporary.cleanup)
        cls.root = Path(cls.temporary.name).resolve()
        evidence.require(cls.root.is_relative_to(generate.PRIVATE_ROOT.resolve()), "unsafe test workspace")
        cls.materials = cls.root / "generated"
        with patch.object(generate, "load_evidence", return_value=(cls.evidence, cls.provenance)):
            cls.manifest = generate.generate(cls.materials)

    def fixture(self, name):
        return shutil.copytree(self.materials, self.root / name)

    def check(self, directory):
        with patch.object(qa, "load_evidence", return_value=(self.evidence, self.provenance)):
            return qa.check_pair(directory)

    def write_manifest(self, directory, manifest):
        (directory / "materials_manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

    def update_output_hash(self, directory, name):
        manifest = evidence.load_json(directory / "materials_manifest.json")
        manifest["output_raw_sha256"][name] = evidence.sha256(directory / name)
        self.write_manifest(directory, manifest)

    def layout(self):
        return {
            "scope": "revised_word_directory_from_measured_render",
            "layout_source_lf_sha256": generate.layout_source_hashes(),
            "page_count": 82,
            "section_pages": {t["number"]: i + 4 for i, t in enumerate(self.topics)},
        }

    def test_validated_source_and_controls_stay_separate(self):
        self.assertEqual(self.provenance["runtime_file_count"], 149)
        self.assertTrue(all(self.provenance["checks"].values()))
        self.assertEqual(len(self.provenance["comparisons"]), 3)
        for report in self.provenance["comparisons"]:
            self.assertTrue(report["optimization_gates_passed"])
            self.assertEqual(len(report["checks"]), 7)
            self.assertFalse(report["all_controls_meet_targets"])
        self.assertFalse(self.analysis["new_performance_run"])
        self.assertFalse(self.analysis["independent_blind"])
        self.assertFalse(self.analysis["external_benchmark_executed"])

    def test_analysis_rejects_unbound_source(self):
        for value in (None, "0" * 64):
            provenance = deepcopy(self.provenance)
            key = (evidence.BASE / "complex-policies.json").as_posix()
            if value is None:
                del provenance["input_raw_sha256"][key]
            else:
                provenance["input_raw_sha256"][key] = value
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "validated provenance"):
                analysis.derive_analysis(provenance)

    def test_analysis_rejects_source_changed_during_read(self):
        key = (evidence.BASE / "complex-policies.json").as_posix()
        with patch.object(analysis, "sha256", side_effect=[self.provenance["input_raw_sha256"][key], "0" * 64]):
            with self.assertRaisesRegex(ValueError, "changed during reading"):
                analysis.derive_analysis(self.provenance)

    def test_policy_counts_and_quality_cost_are_not_speedup(self):
        delta = self.analysis["delta_vs_similarity"]
        self.assertEqual(delta["full_evidence_requests"], 4)
        self.assertEqual(delta["correct_no_answer_requests"], 23)
        self.assertGreater(delta["p95_first_round_ms"], 0)
        self.assertAlmostEqual(delta["p95_first_round_ms"], 28.206, places=3)
        self.assertLess(delta["distractor_rate_percentage_points"], 0)
        self.assertEqual(self.analysis["policies"]["content_coverage_v1"]["full_evidence_correct"], 59)

    def test_metric_example_is_explicitly_illustrative(self):
        example = analysis.worked_example()
        self.assertEqual(example["recall_at_3"], 1)
        self.assertTrue(example["full_evidence"])
        self.assertAlmostEqual(example["precision_at_3"], 2 / 3)
        self.assertAlmostEqual(example["distractor_return_rate"], 1 / 3)
        self.assertAlmostEqual(example["ndcg_at_3"], .9197207891481876)
        self.assertIn("not_an_experiment", example["scope"])

    def test_wilson_uncertainty_and_assumptions(self):
        lower, upper = analysis.wilson(24, 24)
        self.assertAlmostEqual(lower, .8620237953250198)
        self.assertAlmostEqual(upper, 1)
        self.assertLess(analysis.wilson(6, 6)[0], lower)
        self.assertFalse(self.analysis["wilson_assumptions_met"])
        for values in ((True, 1), (1, False), (0, 0), (2, 1), (-1, 1), (1.0, 2)):
            with self.subTest(values=values), self.assertRaises(ValueError):
                analysis.wilson(*values)

    def test_original_technical_chapters_are_preserved(self):
        original = content.build_topics(self.evidence)
        current = {t["number"]: t for t in self.topics}
        self.assertEqual(len(original), 67)
        for t in original:
            if t["number"] not in {"11.3", "18.2"}:
                self.assertEqual(current[t["number"]], t)
        self.assertEqual({t["diagram"] for t in self.topics if t["diagram"]}, set(DIAGRAMS))
        self.assertEqual(len(DIAGRAMS), 25)
        self.assertEqual(content.build_topics(self.evidence), original)

    def test_all_four_scoring_topics_have_detailed_sections(self):
        self.assertEqual(len(CHAPTERS), 18)
        self.assertEqual(len(self.topics), 79)
        self.assertEqual(len({t["number"] for t in self.topics}), 79)
        self.assertEqual(len(FOCUS), 4)
        required = {f"{chapter}.{section}" for chapter, sections in (
            (11, range(4, 9)), (12, range(7, 11)), (13, range(7, 10)),
        ) for section in sections}
        for t in self.topics:
            if t["number"] in required:
                with self.subTest(section=t["number"]):
                    detail = sum(len(p) for _, p in t["paragraphs"])
                    if t["table"]:
                        detail += sum(len(str(cell)) for row in t["table"][1] for cell in row)
                    self.assertGreaterEqual(len(t["paragraphs"]), 3)
                    self.assertGreater(detail, 400)
                    self.assertTrue(t["table"] or t["code"])
        self.assertTrue(required <= {t["number"] for t in self.topics})

    def test_twenty_slide_narrative_and_cross_references(self):
        self.assertEqual(len(self.slides), 20)
        sections = {t["number"] for t in self.topics}
        for s in self.slides:
            self.assertTrue(s["sections"] and set(s["sections"]) <= sections)
            self.assertTrue(s["speaker"] and s["source"])
            if s["kind"] == "points":
                self.assertEqual(len(s["points"]), 3)
            if s["table"]:
                headers, rows = s["table"]
                self.assertLessEqual(len(rows), 5)
                self.assertTrue(all(len(row) == len(headers) for row in rows))
        self.assertEqual(len({s["diagram"] for s in self.slides if s["diagram"]}), 9)

    def test_page_map_binds_layout_source(self):
        layout = self.layout()
        generate.validate_word_layout(layout, self.topics)
        key = next(iter(layout["layout_source_lf_sha256"]))
        layout["layout_source_lf_sha256"][key] = "0" * 64
        with self.assertRaisesRegex(ValueError, "page map"):
            generate.validate_word_layout(layout, self.topics)

    def test_page_map_rejects_missing_invalid_or_unsorted_pages(self):
        for value in (True, 0, 83, "4", 6):
            layout = self.layout()
            layout["section_pages"]["1.1"] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                generate.validate_word_layout(layout, self.topics)
        layout = self.layout()
        del layout["section_pages"]["1.1"]
        with self.assertRaises(ValueError):
            generate.validate_word_layout(layout, self.topics)

    def test_existing_outputs_and_outside_workspace_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "overwrite"):
            generate.generate(self.materials)
        outside = ROOT.parent / "must-not-create-revised-materials"
        with self.assertRaisesRegex(ValueError, "stay in"):
            generate.generate(outside)
        self.assertFalse(outside.exists())

    def test_pair_passes_strict_checks(self):
        result = self.check(self.materials)
        self.assertTrue(result["passed"])
        self.assertTrue(all(result["checks"].values()))
        self.assertEqual(result["word"]["native_tables"], 49)
        self.assertEqual(result["presentation"]["slides"], 20)
        self.assertEqual(result["presentation"]["editable_diagrams"], 10)
        self.assertGreaterEqual(result["presentation"]["minimum_table_font_pt"], 16)
        self.assertFalse(result["competition_ready"])

    def test_changed_office_bytes_are_rejected(self):
        directory = self.fixture("office-drift")
        with (directory / generate.DOCX_NAME).open("ab") as stream:
            stream.write(b"drift")
        with self.assertRaisesRegex(ValueError, "material output changed"):
            self.check(directory)

    def test_missing_hash_entry_is_rejected(self):
        directory = self.fixture("missing-hash")
        manifest = evidence.load_json(directory / "materials_manifest.json")
        del manifest["output_raw_sha256"][generate.PPTX_NAME]
        self.write_manifest(directory, manifest)
        with self.assertRaisesRegex(ValueError, "inventory"):
            self.check(directory)

    def test_unsupported_claim_is_rejected(self):
        directory = self.fixture("false-claims")
        original = evidence.load_json(directory / "materials_manifest.json")
        for key in qa.UNSUPPORTED_CLAIMS:
            manifest = deepcopy(original)
            manifest[key] = True
            self.write_manifest(directory, manifest)
            with self.subTest(key=key), self.assertRaisesRegex(ValueError, "unsupported"):
                self.check(directory)

    def test_provenance_and_derived_values_are_revalidated(self):
        for index, (name, key) in enumerate((
            ("evidence_provenance.json", "runtime_file_count"), ("analysis.json", "new_performance_run"),
        )):
            directory = self.fixture(f"evidence-drift-{index}")
            value = evidence.load_json(directory / name)
            value[key] = True
            (directory / name).write_text(json.dumps(value), encoding="utf-8")
            self.update_output_hash(directory, name)
            with self.assertRaisesRegex(ValueError, "provenance changed|analysis changed"):
                self.check(directory)

    def test_tool_snapshot_is_required(self):
        directory = self.fixture("tool-drift")
        manifest = evidence.load_json(directory / "materials_manifest.json")
        key = next(iter(manifest["document_tool_lf_sha256"]))
        manifest["document_tool_lf_sha256"][key] = "0" * 64
        self.write_manifest(directory, manifest)
        with self.assertRaisesRegex(ValueError, "tool snapshot"):
            self.check(directory)

    def test_detailed_word_text_cannot_be_removed_and_rehashed(self):
        directory = self.fixture("word-detail-removed")
        path = directory / generate.DOCX_NAME
        doc = Document(path)
        body = next(t for t in self.topics if t["number"] == "11.6")["paragraphs"][0][1]
        paragraph = next(p for p in doc.paragraphs if body in p.text)
        paragraph.text = "Removed technical discussion."
        doc.save(path)
        self.update_output_hash(directory, generate.DOCX_NAME)
        with self.assertRaisesRegex(ValueError, "missing Word detail"):
            self.check(directory)

    def test_sdk_risk_must_remain_visible_not_just_in_notes(self):
        directory = self.fixture("sdk_notes_only")
        path = directory / generate.PPTX_NAME
        prs = Presentation(path)
        boundary = next(s for s in prs.slides[7].shapes if s.name == "visible_boundary")
        notes = prs.slides[7].notes_slide.notes_text_frame
        notes.text += "\n" + boundary.text
        boundary.text = "Boundary removed."
        prs.save(path)
        self.update_output_hash(directory, generate.PPTX_NAME)
        with self.assertRaisesRegex(ValueError, "visible slide 8"):
            self.check(directory)

    def test_small_ppt_table_font_is_rejected(self):
        directory = self.fixture("small-table-font")
        path = directory / generate.PPTX_NAME
        prs = Presentation(path)
        table = next(s.table for s in prs.slides[8].shapes if s.has_table)
        table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        prs.save(path)
        self.update_output_hash(directory, generate.PPTX_NAME)
        with self.assertRaisesRegex(ValueError, "table text too small"):
            self.check(directory)

    def test_pdf_and_zip_are_not_delivered(self):
        for suffix in (".pdf", ".zip"):
            directory = self.fixture("extra" + suffix)
            (directory / ("unexpected" + suffix)).write_bytes(b"test")
            with self.subTest(suffix=suffix), self.assertRaisesRegex(ValueError, "PDF or ZIP"):
                self.check(directory)

    def test_render_record_binds_source_and_all_output_bytes(self):
        directory = self.root / "render-record-test"
        directory.mkdir()
        output = directory / "test.png"
        output.write_bytes(b"unit test, not a real render")
        record = {
            "scope": "read_only_wps_render_for_private_layout_review", "kind": "ppt",
            "source_unchanged": True, "source_raw_sha256": evidence.sha256(self.materials / generate.PPTX_NAME),
            "output_raw_sha256": {"test.png": evidence.sha256(output)},
        }
        (directory / "render_manifest.json").write_text(json.dumps(record), encoding="utf-8")
        self.assertEqual(qa.check_render_manifest(self.materials / generate.PPTX_NAME, directory, "ppt"), record)
        output.write_bytes(b"changed output")
        with self.assertRaisesRegex(ValueError, "render output changed"):
            qa.check_render_manifest(self.materials / generate.PPTX_NAME, directory, "ppt")
        with self.assertRaisesRegex(ValueError, "bound to its Office source"):
            qa.check_render_manifest(self.materials / generate.DOCX_NAME, directory, "ppt")


if __name__ == "__main__":
    unittest.main()
