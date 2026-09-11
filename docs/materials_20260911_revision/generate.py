"""Generate the twenty-slide report and expanded Word without replacing prior files."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as WI, Pt as WP
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PI, Pt as PP

from docs.materials_20260911 import generate as base
from docs.materials_20260911.content import DATE, STATUS, TITLE
from docs.materials_20260911.diagrams import DIAGRAMS
from docs.materials_20260911.evidence import lf_sha256, load_evidence, load_json, require, sha256
from docs.materials_20260911_revision.analysis import derive_analysis
from docs.materials_20260911_revision.slide_content import build_slides
from docs.materials_20260911_revision.word_content import CHAPTERS, FOCUS, build_topics

DOCX_NAME = "OS_Agent_技术方案_评分重点详述版_20260911.docx"
PPTX_NAME = "OS_Agent_项目报告_20页精简版_20260911.pptx"
DEFAULT_OUTPUT = ROOT / "deliverables/submission-materials-20260911-revised"
PRIVATE_ROOT = Path.home() / "TiaobeiLab"
SCOPE = "twenty_slide_report_and_expanded_word_not_competition_acceptance"


def tool_hashes():
    folders = (ROOT / "docs/materials_20260911", Path(__file__).parent)
    return {path.relative_to(ROOT).as_posix(): lf_sha256(path)
            for folder in folders for path in sorted(folder.iterdir())
            if path.suffix in (".py", ".ps1")}


def layout_source_hashes():
    paths = [
        "docs/materials_20260911/content.py", "docs/materials_20260911/diagrams.py",
        "docs/materials_20260911/generate.py", "docs/materials_20260911_revision/analysis.py",
        "docs/materials_20260911_revision/word_content.py",
        "docs/materials_20260911_revision/generate.py",
    ]
    return {path: lf_sha256(ROOT / path) for path in paths}


def validate_word_layout(layout, topics):
    require(layout.get("scope") == "revised_word_directory_from_measured_render"
            and layout.get("layout_source_lf_sha256") == layout_source_hashes()
            and set(layout.get("section_pages", {})) == {t["number"] for t in topics}
            and type(layout.get("page_count")) is int and layout["page_count"] >= 4,
            "Word page map does not match revised source")
    pages = [layout["section_pages"][t["number"]] for t in topics]
    require(all(type(page) is int and 4 <= page <= layout["page_count"] for page in pages)
            and pages == sorted(pages), "Word page map contains invalid or unordered pages")


def word_cover(doc, topics, e):
    base.word_paragraph(doc, "挑战杯 · 揭榜挂帅  |  麒麟赛题", size=12, bold=True, color="teal",
                        before=24, after=25, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.word_paragraph(doc, "OS Agent\n多源融合记忆优化系统", size=27, bold=True,
                        after=20, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.3)
    base.word_paragraph(doc, "技术方案", size=23, bold=True, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.word_paragraph(doc, "评分重点详述版 · 含用户手册与效果验证报告", size=12,
                        color="muted", after=19, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.word_paragraph(doc, STATUS, size=13, bold=True, color="teal",
                        after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.word_paragraph(doc, f"{DATE}  |  匿名评审稿  |  配套20页项目报告",
                        size=10, color="muted", after=19, align=WD_ALIGN_PARAGRAPH.CENTER)
    base.word_table(doc, (["重点", "详细位置", "论证内容"], FOCUS), "focus", size=10)
    base.word_paragraph(doc, f"完整范围：18章、{len(topics)}节、25张方案框图。"
                        f"沿用已核验公开合成证据：{e['tests']}项工程测试无失败、{e['skipped']}项平台跳过。",
                        size=10.5, before=12, after=8)
    base.word_paragraph(doc, "30分技术创新性与25分性能可信度是反馈强调的论证重点，非项目自评分。"
                        "HashEmbedding不代表真实SDK；真实SDK为待复测风险项，当前源码V11、独立Agent/Judge及授权案例另验。",
                        size=10, color="amber", after=9)
    base.word_paragraph(doc, "仅交付可编辑Word与PPTX；PDF由负责人转换。本次不打包、不上传，旧版材料保留。",
                        size=9.5, color="muted", align=WD_ALIGN_PARAGRAPH.CENTER)


def word_navigation(doc, topics, page_map):
    doc.add_page_break()
    base.word_paragraph(doc, "目录与重点阅读路径", size=22, bold=True, after=14)
    base.word_paragraph(doc, "方案与实现见第2–10章；相关工作与技术创新性详见第11章；"
                        "评测口径与性能可信度详见第12–13章；第14–18章保留案例、用户手册、接口及验收证据。",
                        size=11, after=17)
    for chapter, title in CHAPTERS.items():
        target = next(t for t in topics if t["number"].split(".")[0] == chapter)
        p = base.word_paragraph(doc, "", size=11, after=7)
        p.paragraph_format.tab_stops.add_tab_stop(WI(base.BODY_WIDTH - .35))
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "sec_" + target["number"].replace(".", "_"))
        run, props, color, text = (OxmlElement(tag) for tag in ("w:r", "w:rPr", "w:color", "w:t"))
        color.set(qn("w:val"), base.C["teal"])
        props.append(color)
        run.append(props)
        text.text = f"第{chapter}章  {title}"
        run.append(text)
        link.append(run)
        p._p.append(link)
        p.add_run("\t")
        base.field(p, "PAGEREF sec_" + target["number"].replace(".", "_") + " \\h",
                   page_map.get(target["number"], 4), 11)
    base.word_paragraph(doc, "目录使用可更新页码域；本版按实际渲染分页回填。"
                        "评分重点：11.4–11.8机制与创新，12.7–12.10真值/公式/不确定性，13.7–13.9可信度与风险。",
                        size=9.5, color="muted", before=15)
    doc.add_page_break()
    base.word_paragraph(doc, "方案框图索引", size=22, bold=True, after=12)
    base.word_paragraph(doc, "本技术方案保留25张高清框图。配套20页PPT选择关键框图，采用可编辑形状、文字和连接线。",
                        size=11, after=16)
    count = 0
    for t in topics:
        if not t["diagram"]:
            continue
        count += 1
        p = base.word_paragraph(doc, f"图{count:02d}  {DIAGRAMS[t['diagram']]['title']}",
                                size=10, after=4)
        p.paragraph_format.tab_stops.add_tab_stop(WI(base.BODY_WIDTH - .35))
        p.add_run("\t")
        base.field(p, "PAGEREF sec_" + t["number"].replace(".", "_") + " \\h",
                   page_map.get(t["number"], 4), 10)
    base.word_paragraph(doc, "框图表示实际模块和处理边界；条件失败可提前退出。"
                        "两路线无自动同步，真实模型、授权与正式验收边界在对应章节单列。",
                        size=9.5, color="muted", before=15)


def make_word(path, topics, e, layout=None):
    pages = layout["section_pages"] if layout else {}
    doc = base.setup_word(layout["page_count"] if layout else 1)
    doc.core_properties.subject = "详细技术方案、相关工作、技术创新性、评测口径与性能可信度"
    word_cover(doc, topics, e)
    word_navigation(doc, topics, pages)
    figures, seen = 0, set()
    for i, t in enumerate(topics):
        chapter = t["number"].split(".")[0]
        if chapter not in seen:
            doc.add_page_break()
            doc.add_paragraph(f"第{chapter}章  {CHAPTERS[chapter]}", "Heading 1")
            seen.add(chapter)
        heading = doc.add_paragraph(f"{t['number']}  {t['title']}", "Heading 2")
        if not t["number"].endswith(".1"):
            heading.paragraph_format.space_before = WP(17)
        base.bookmark(heading, "sec_" + t["number"].replace(".", "_"), i + 1)
        base.word_paragraph(doc, t["lead"], size=11.5, color="teal", after=11, line=1.22, keep=True)
        if t["diagram"]:
            figures += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = WP(4)
            p.paragraph_format.keep_with_next = True
            image = p.add_run().add_picture(base.diagram_png(t["diagram"]), width=WI(base.BODY_WIDTH))
            spec = DIAGRAMS[t["diagram"]]
            image._inline.docPr.set("descr", spec["title"] + "；" + " → ".join(
                node["title"] + "：" + node["body"].replace("\n", "，") for node in spec["nodes"]))
            base.word_paragraph(doc, f"图{figures:02d}  {spec['title']}。{spec['note']}",
                                size=9, color="muted", after=10, line=1.15)
        for label, text in t["paragraphs"]:
            p = base.word_paragraph(doc, "", size=12, after=6, line=1.20,
                                    align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            base.word_font(p.add_run(label + "。"), 12, True)
            base.word_font(p.add_run(text), 12)
        if t["table"]:
            base.word_table(doc, t["table"], t["number"], size=9.2 if len(t["table"][1]) > 8 else 9.6)
        if t["code"]:
            p = base.word_paragraph(doc, t["code"], size=9.2, face=base.MONO, line=1.18, after=8)
            base.shade(p._p.get_or_add_pPr(), "blue_pale")
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = WI(.08)
            p.paragraph_format.keep_together = True
        if t["warning"]:
            base.word_paragraph(doc, t["warning"], size=9.5, bold=True, color="amber", after=8)
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        source = base.word_paragraph(doc, "依据：" + t["source"], size=8.2, color="muted", before=3, after=0)
        source.paragraph_format.keep_together = True
    doc.save(path)
    return {"expected_pages": layout["page_count"] if layout else None,
            "section_pages": pages, "page_map_from_render": bool(layout),
            "chapters": len(CHAPTERS), "topics": len(topics),
            "figures": figures, "native_tables": len(doc.tables),
            "scoring_focus": [dict(zip(("topic", "sections", "detail"), row)) for row in FOCUS]}


TABLE_RATIOS = {
    9: (.24, .30, .46), 10: (.26, .32, .42), 13: (.22, .15, .34, .29),
    14: (.19, .36, .45), 16: (.31, .23, .20, .26), 17: (.31, .23, .23, .23),
    18: (.27, .34, .39),
}


def native_table(slide, spec, number, *, y=2.02, max_h=4.13):
    headers, rows = spec
    ratios = TABLE_RATIOS[number]
    require(len(ratios) == len(headers), "table widths do not match columns")
    size = 18
    while size >= 16:
        wrapped, heights = [], []
        for index, values in enumerate([headers, *rows]):
            cells = [base.wrap(str(value), 12 * ratios[j] * 72 - 18, size, index == 0)
                     for j, value in enumerate(values)]
            wrapped.append(cells)
            heights.append(max(map(len, cells)) * size * 1.18 / 72 + .19)
        if sum(heights) <= max_h:
            break
        size -= .5
    require(size >= 16 and sum(heights) <= max_h, f"slide {number} table is too dense")
    graphic = slide.shapes.add_table(len(rows) + 1, len(headers), PI(.67), PI(y), PI(12), PI(sum(heights)))
    graphic.name = f"editable_table:{number}"
    table = graphic.table
    for j, ratio in enumerate(ratios):
        table.columns[j].width = PI(12 * ratio)
    for index, cells in enumerate(wrapped):
        table.rows[index].height = PI(heights[index])
        for j, lines in enumerate(cells):
            cell = table.cell(index, j)
            cell.margin_left = cell.margin_right = PI(.11)
            cell.margin_top = cell.margin_bottom = PI(.07)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                base.C["teal" if index == 0 else "pale" if index % 2 else "paper"])
            cell.text_frame.clear()
            cell.text_frame.word_wrap = False
            for k, value in enumerate(lines):
                p = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                p.line_spacing = PP(size * 1.18)
                p.space_before = p.space_after = PP(0)
                run = p.add_run()
                run.text = value
                base.ppt_font(run, size, index == 0, "paper" if index == 0 else "ink")
            for side in ("lnL", "lnR", "lnT", "lnB"):
                border = base.PX("a:" + side)
                border.set("w", "6350")
                solid, color = base.PX("a:solidFill"), base.PX("a:srgbClr")
                color.set("val", "FFFFFF")
                solid.append(color)
                border.append(solid)
                cell._tc.get_or_add_tcPr().append(border)
    return size


def timing_strip(slide):
    labels = [
        ("开始计时", "调用 search 之前"),
        ("锁、编码与候选", "真实 provider 成本"),
        ("重排、正文与返回", "包含序列化"),
        ("结束计时", "nearest-rank 统计"),
    ]
    for i, (title, text) in enumerate(labels):
        x = .69 + 3.06 * i
        if i < 3:
            base.connector(slide, (x + 2.70, 2.44), (x + 3.03, 2.44))
        box = base.rect(slide, x, 1.99, 2.70, .91, "blue_pale" if i in (1, 2) else "pale",
                        "blue" if i in (1, 2) else "teal")
        box.name = f"diagram:timing_boundary:node:{i + 1}"
        base.textbox(slide, title, x + .08, 2.13, 2.54, .31, 17, bold=True,
                     align=PP_ALIGN.CENTER, color="blue" if i in (1, 2) else "teal")
        base.textbox(slide, text, x + .07, 2.56, 2.56, .24, 13.5, align=PP_ALIGN.CENTER)


def ppt_cover(prs, e):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(base.C["paper"])
    base.rect(slide, .77, .65, .07, .33, "teal")
    base.textbox(slide, "挑战杯 · 揭榜挂帅  |  麒麟赛题  |  匿名项目报告", 1.02, .68, 11.5, .34,
                 15, bold=True, color="teal")
    base.textbox(slide, "OS Agent\n多源融合记忆优化系统", .82, 1.36, 11.74, 1.68,
                 42, bold=True, spacing=1.22)
    base.textbox(slide, "总体方案 · 技术创新 · 评测口径 · 性能可信度", .88, 3.40, 11.8, .45,
                 23, color="muted")
    for i, (number, label, body) in enumerate((
        (str(e["tests"]), "项工程测试无失败", "1项平台条件跳过"),
        ("168", "个不同复杂请求", "可见开发30题另列"),
        (f"{e['frozen_p95']:.3f} ms", "冻结检索 P95", "80查询 / HashEmbedding"),
    )):
        x = .89 + 4.06 * i
        tone = "blue" if i == 1 else "teal"
        base.rect(slide, x, 4.31, 3.64, .035, tone)
        base.textbox(slide, number, x, 4.61, 3.70, .55, 31, bold=True, color=tone)
        base.textbox(slide, label, x, 5.32, 3.70, .36, 18.5, bold=True)
        base.textbox(slide, body, x, 5.83, 3.70, .30, 15, color="muted")
    base.textbox(slide, STATUS + "  |  " + DATE, .89, 6.47, 11.80, .36, 17, bold=True, color="teal")
    base.textbox(slide, "公开合成回归；真实SDK为待复测风险项，当前V11与独立Agent/Judge另验。",
                 .89, 7.03, 11.50, .25, 11.5, color="amber")
    base.textbox(slide, "01 / 20", 12.06, 7.27, .61, .20, 9, align=PP_ALIGN.RIGHT, color="muted")
    return slide


def make_ppt(path, specs, e):
    require(len(specs) == 20, "presentation must contain exactly twenty slides")
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(base.SLIDE_WIDTH), PI(base.SLIDE_HEIGHT)
    prs.core_properties.title = TITLE + " 项目报告（20页精简版）"
    prs.core_properties.subject = "总体方案、技术创新性、评测口径与性能可信度"
    prs.core_properties.author = prs.core_properties.last_modified_by = ""
    prs.core_properties.keywords = "OS Agent; anonymous; public regression"
    prs.core_properties.comments = ""
    prs.core_properties.created = prs.core_properties.modified = base.GENERATED_AT
    diagram_count = table_count = 0
    minimum_table_size = 100
    for index, s in enumerate(specs, 1):
        if s["kind"] == "cover":
            slide = ppt_cover(prs, e)
        else:
            slide = base.slide_base(prs, index, 20, s["title"], s["chapter"], s["lead"], s["source"])
            if s["kind"] == "diagram":
                base.ppt_diagram(slide, s["diagram"], y=1.87)
                diagram_count += 1
            elif s["kind"] in ("table", "timing"):
                if s["kind"] == "timing":
                    timing_strip(slide)
                    diagram_count += 1
                size = native_table(slide, s["table"], index,
                                    y=3.20 if s["kind"] == "timing" else 2.02,
                                    max_h=2.80 if s["kind"] == "timing" else 4.13)
                minimum_table_size = min(minimum_table_size, size)
                table_count += 1
            else:
                require(len(s["points"]) == 3, "point slides require three concise arguments")
                for i, (title, text) in enumerate(s["points"]):
                    y = 2.06 + i * 1.38
                    base.textbox(slide, f"0{i + 1}", .78, y, .72, .43, 25, bold=True, color="teal")
                    base.textbox(slide, title, 1.66, y, 10.78, .39, 21, bold=True)
                    base.textbox(slide, text, 1.68, y + .51, 10.70, .72,
                                 17, color="muted", spacing=1.16)
            if s["note"]:
                base.textbox(slide, s["note"], .77, 6.35, 11.85, .70, 14,
                             minimum=13.5, color="amber", spacing=1.16, name="visible_boundary")
        slide.notes_slide.notes_text_frame.text = (
            f"{index:02d} {s['title']}\n{s['lead']}\n\n{s['speaker']}\n\n"
            f"配套Word：{', '.join(s['sections'])}节。\n依据：{s['source']}\n"
            "测量边界：既有Windows/HashEmbedding公开合成回归；不构成真实SDK或赛事验收。"
        )
    prs.save(path)
    return {"slides": len(prs.slides), "editable_diagrams": diagram_count,
            "native_tables": table_count, "minimum_table_font_pt": minimum_table_size,
            "speaker_notes": True, "body_scoring_topics_visible": True}


def generate(output=DEFAULT_OUTPUT, word_layout=None):
    output = Path(output).resolve()
    require(output.is_relative_to(ROOT) or output.is_relative_to(PRIVATE_ROOT.resolve()),
            "output must stay in the project or private lab")
    require(not output.exists(), "refusing to overwrite an existing material directory")
    e, provenance = load_evidence()
    analysis = derive_analysis(provenance)
    topics, slides = build_topics(e, analysis), build_slides(e, analysis)
    require(len(topics) == 79 and len(slides) == 20, "revision scope changed")
    sections = {t["number"] for t in topics}
    require(len(sections) == len(topics)
            and all(s["sections"] and set(s["sections"]) <= sections for s in slides),
            "presentation references a missing or duplicate Word section")
    require({t["diagram"] for t in topics if t["diagram"]} == set(DIAGRAMS),
            "all detailed Word architecture diagrams must be retained")
    layout = load_json(word_layout) if word_layout else None
    if layout:
        validate_word_layout(layout, topics)
    output.mkdir(parents=True)
    word = make_word(output / DOCX_NAME, topics, e, layout)
    presentation = make_ppt(output / PPTX_NAME, slides, e)
    base.write_json(output / "evidence_provenance.json", provenance)
    base.write_json(output / "analysis.json", analysis)
    base.write_json(output / "content_index.json", {
        "word": [{"section": t["number"], "title": t["title"], "diagram": t["diagram"],
                  "expected_page": word["section_pages"].get(t["number"])} for t in topics],
        "presentation": [{"slide": i, "title": s["title"], "word_sections": s["sections"],
                          "diagram": s["diagram"] or ("timing_boundary" if s["kind"] == "timing" else None)}
                         for i, s in enumerate(slides, 1)],
        "scoring_focus": word["scoring_focus"],
    })
    manifest = {
        "schema_version": 1, "scope": SCOPE, "date": DATE, "display_status": STATUS,
        "supersedes_for_current_review": "submission-materials-20260911",
        "word": word, "presentation": presentation,
        "document_tool_lf_sha256": tool_hashes(),
        "layout_source_lf_sha256": layout_source_hashes(),
        "output_raw_sha256": {path.name: sha256(path) for path in sorted(output.iterdir()) if path.is_file()},
        "source_validation_passed": all(provenance["checks"].values()),
        "visual_render_verified": False,
        "new_performance_run": False, "external_benchmark_executed": False,
        "native_latency_passed": False, "current_v11_verified": False,
        "independent_blind": False, "competition_ready": False,
        "pdf_delivered": False, "packaged": False, "uploaded": False,
    }
    base.write_json(output / "materials_manifest.json", manifest)
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--word-layout", type=Path)
    args = parser.parse_args(argv)
    manifest = generate(args.output_dir, args.word_layout)
    print(json.dumps({key: manifest[key] for key in ("word", "presentation", "source_validation_passed")},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
