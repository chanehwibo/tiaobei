"""Validate the revised Office pair and its private read-only rendering."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docx import Document
from PIL import Image, ImageStat
from pptx import Presentation

from docs.materials_20260911.content import STATUS
from docs.materials_20260911.diagrams import DIAGRAMS
from docs.materials_20260911.evidence import load_evidence, load_json, require, sha256
from docs.materials_20260911.generate import write_json
from docs.materials_20260911.qa import contacts, normalized, package_text, verify_render_equivalent
from docs.materials_20260911_revision.analysis import derive_analysis
from docs.materials_20260911_revision.generate import (
    DOCX_NAME, PPTX_NAME, PRIVATE_ROOT, SCOPE, layout_source_hashes, tool_hashes, validate_word_layout,
)
from docs.materials_20260911_revision.slide_content import build_slides
from docs.materials_20260911_revision.word_content import CHAPTERS, FOCUS, build_topics
from scripts.check_deliverables import inspect_ooxml

UNSUPPORTED_CLAIMS = (
    "native_latency_passed", "current_v11_verified", "independent_blind", "competition_ready",
    "new_performance_run", "external_benchmark_executed", "pdf_delivered", "packaged", "uploaded",
)


def visible_text(slide):
    values = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            values.append(shape.text)
        if shape.has_table:
            values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return normalized("".join(values))


def check_pair(directory, *, check_tool_hashes=True):
    directory = Path(directory).resolve()
    evidence, provenance = load_evidence()
    analysis = derive_analysis(provenance)
    topics, specs = build_topics(evidence, analysis), build_slides(evidence, analysis)
    manifest = load_json(directory / "materials_manifest.json")
    require(manifest.get("scope") == SCOPE and manifest.get("display_status") == STATUS,
            "material status or scope changed")
    require(all(manifest.get(key) is False for key in UNSUPPORTED_CLAIMS),
            "unsupported acceptance, measurement or delivery claim")
    require(manifest.get("source_validation_passed") is True, "source validation not recorded")
    require(set(manifest["output_raw_sha256"]) == {
        DOCX_NAME, PPTX_NAME, "evidence_provenance.json", "analysis.json", "content_index.json",
    }, "material output inventory changed")
    require(load_json(directory / "evidence_provenance.json") == provenance,
            "document provenance changed")
    require(load_json(directory / "analysis.json") == analysis, "derived analysis changed")
    for name, expected in manifest["output_raw_sha256"].items():
        path = (directory / name).resolve()
        require(path.is_relative_to(directory) and path.is_file() and sha256(path) == expected,
                "material output changed: " + name)
    if check_tool_hashes:
        require(manifest["document_tool_lf_sha256"] == tool_hashes(), "document tool snapshot changed")
        require(manifest["layout_source_lf_sha256"] == layout_source_hashes(), "layout source changed")
    if manifest["word"].get("page_map_from_render"):
        validate_word_layout({
            "scope": "revised_word_directory_from_measured_render",
            "layout_source_lf_sha256": manifest["layout_source_lf_sha256"],
            "page_count": manifest["word"]["expected_pages"],
            "section_pages": manifest["word"]["section_pages"],
        }, topics)
    require(not any(p.suffix.lower() in (".pdf", ".zip") for p in directory.rglob("*")),
            "this delivery must not contain PDF or ZIP")
    require(sorted(p.name for p in directory.glob("*.docx")) == [DOCX_NAME]
            and sorted(p.name for p in directory.glob("*.pptx")) == [PPTX_NAME],
            "expected exactly one Word and one project report")
    for name in (DOCX_NAME, PPTX_NAME):
        result = inspect_ooxml(directory / name, expected_slide_count=20)
        require(result["passed"], "invalid Office structure: " + repr(result["issues"]))
    word_text, ppt_text = (package_text(directory / name) for name in (DOCX_NAME, PPTX_NAME))
    for text in (word_text, ppt_text):
        for required in (
            STATUS, "待复测风险项", "HashEmbedding", "48.145", "13,390.615", "MemGPT",
            "Mem0", "MemoryOS", "LongMemEval", "nearest-rank", "46.15%", "31.71%",
            "2026年9月15日24:00", "5–10", "200M", "技术创新性", "性能可信度",
        ):
            require(normalized(required) in text, "missing disclosure: " + required)
    for t in topics:
        require(normalized(t["title"]) in word_text, "missing Word section: " + t["number"])
        for label, paragraph in t["paragraphs"]:
            require(normalized(label + "。" + paragraph) in word_text,
                    "missing Word detail: " + t["number"])
        for value in (t["lead"], t["warning"], t["code"], "依据：" + t["source"]):
            require(normalized(value) in word_text, "missing Word supporting detail: " + t["number"])
        if t["table"]:
            for row in (t["table"][0], *t["table"][1]):
                for value in row:
                    require(normalized(value) in word_text, "missing Word cell: " + t["number"])
    doc = Document(directory / DOCX_NAME)
    require(len(doc.inline_shapes) == len(DIAGRAMS) == manifest["word"]["figures"] == 25,
            "Word diagram count changed")
    require(len(doc.tables) == manifest["word"]["native_tables"] == 49, "Word table count changed")
    require(len([p for p in doc.paragraphs if p.style.name == "Heading 2"]) == 79
            and manifest["word"]["topics"] == 79
            and manifest["word"]["chapters"] == len(CHAPTERS) == 18, "Word section count changed")
    require(manifest["word"]["scoring_focus"] == [
        dict(zip(("topic", "sections", "detail"), row)) for row in FOCUS
    ], "scoring focus index changed")
    prs = Presentation(directory / PPTX_NAME)
    require(len(prs.slides) == len(specs) == manifest["presentation"]["slides"] == 20,
            "presentation must contain twenty slides")
    diagrams, tables, sizes = set(), 0, []
    for index, (slide, spec) in enumerate(zip(prs.slides, specs), 1):
        text, notes = visible_text(slide), normalized(slide.notes_slide.notes_text_frame.text)
        for value in (spec["title"], spec["lead"], spec["note"]):
            require(normalized(value) in text, f"missing visible slide {index} content")
        for title, body in spec["points"]:
            require(normalized(title) in text and normalized(body) in text,
                    f"missing visible point on slide {index}")
        if spec["table"]:
            for row in (spec["table"][0], *spec["table"][1]):
                for value in row:
                    require(normalized(value) in text, f"missing visible table cell on slide {index}")
        require(normalized(spec["speaker"]) in notes and normalized(spec["source"]) in notes
                and all(normalized(section) in notes for section in spec["sections"]),
                f"missing curated speaker notes on slide {index}")
        require(set(spec["sections"]) <= {t["number"] for t in topics},
                f"invalid Word cross-reference on slide {index}")
        for shape in slide.shapes:
            require(min(shape.left, shape.top, shape.width, shape.height) >= 0
                    and shape.left + shape.width <= prs.slide_width + 8
                    and shape.top + shape.height <= prs.slide_height + 8,
                    f"shape outside slide {index}: {shape.name}")
            require(shape.shape_type != 13, "PPT block diagrams must remain editable")
            if shape.name.startswith("diagram:"):
                diagrams.add(shape.name.split(":")[1])
            if shape.has_table:
                tables += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                require(run.font.size is not None, "table font not explicit")
                                sizes.append(run.font.size.pt)
    expected_diagrams = {s["diagram"] for s in specs if s["diagram"]} | {"timing_boundary"}
    require(diagrams == expected_diagrams and len(diagrams) == manifest["presentation"]["editable_diagrams"] == 10
            and tables == manifest["presentation"]["native_tables"] == 7, "PPT diagram/table count changed")
    require(sizes and min(sizes) >= 16 and min(sizes) == manifest["presentation"]["minimum_table_font_pt"],
            "PPT table text too small or manifest inconsistent")
    full_body = "".join(visible_text(s) for s in prs.slides)
    for phrase in ("相关工作对比", "评测口径说明", "技术创新性", "性能可信度", "待复测风险项",
                   "13,390.615", "31.71%", "不是普遍加速", "合成"):
        require(normalized(phrase) in full_body, "scoring topic or caveat only appears in notes: " + phrase)
    index = load_json(directory / "content_index.json")
    require(index["word"] == [
        {"section": t["number"], "title": t["title"], "diagram": t["diagram"],
         "expected_page": manifest["word"]["section_pages"].get(t["number"])} for t in topics
    ] and index["presentation"] == [
        {"slide": i, "title": s["title"], "word_sections": s["sections"],
         "diagram": s["diagram"] or ("timing_boundary" if s["kind"] == "timing" else None)}
        for i, s in enumerate(specs, 1)
    ] and index["scoring_focus"] == manifest["word"]["scoring_focus"], "content index changed")
    return {
        "scope": "revised_ooxml_content_anonymity_editability_and_evidence",
        "passed": True, "runtime_file_count": provenance["runtime_file_count"],
        "word": manifest["word"], "presentation": manifest["presentation"],
        "office_sha256": {name: sha256(directory / name) for name in (DOCX_NAME, PPTX_NAME)},
        "checks": {
            "source_and_evidence_revalidated": True, "output_hashes_match": True,
            "tool_hashes_checked": check_tool_hashes, "derived_analysis_recomputed": True,
            "all_word_details_and_tables_present": True, "twenty_slides_with_curated_notes": True,
            "scoring_topics_and_risks_visible": True, "word_cross_references_valid": True,
            "identity_and_obsolete_status_scan_passed": True, "anonymous_metadata": True,
            "native_diagrams_and_tables_verified": True, "shapes_inside_canvas": True,
            "no_delivered_pdf_or_zip": True,
        },
        "visual_review_required": True, "competition_ready": False,
    }


def check_render_manifest(source, directory, kind):
    directory = Path(directory).resolve()
    record = load_json(directory / "render_manifest.json")
    require(record.get("scope") == "read_only_wps_render_for_private_layout_review"
            and record.get("kind") == kind and record.get("source_unchanged") is True
            and record.get("source_raw_sha256") == sha256(source),
            "render is not bound to its Office source")
    files = record["output_raw_sha256"]
    actual = {p.name for p in directory.iterdir() if p.is_file() and p.name != "render_manifest.json"}
    require(files and set(files) == actual, "render inventory changed")
    for name, expected in files.items():
        path = (directory / name).resolve()
        require(path.is_relative_to(directory) and path.is_file() and sha256(path) == expected,
                "render output changed: " + name)
    return record


def rendered_checks(materials, word_pdf, ppt_dir, output, *, check_tool_hashes=True,
                    ppt_render_source=None):
    import fitz

    materials, word_pdf, ppt_dir, output = map(Path, (materials, word_pdf, ppt_dir, output))
    output = output.resolve()
    require(output.is_relative_to(PRIVATE_ROOT.resolve()), "QA rasters must stay in the private lab")
    require(not output.exists(), "refusing to overwrite an existing render review")
    structural = check_pair(materials, check_tool_hashes=check_tool_hashes)
    equivalent = (verify_render_equivalent(materials / PPTX_NAME, ppt_render_source)
                  if ppt_render_source else None)
    word_record = check_render_manifest(materials / DOCX_NAME, word_pdf.parent, "word")
    ppt_record = check_render_manifest(ppt_render_source or materials / PPTX_NAME, ppt_dir, "ppt")
    e, provenance = load_evidence()
    topics = build_topics(e, derive_analysis(provenance))
    output.mkdir(parents=True)
    problems, word_pages, word_images, page_texts, section_pages = [], [], [], [], {}
    with fitz.open(word_pdf) as pdf:
        page_count = len(pdf)
        expected = structural["word"]["expected_pages"]
        if (expected is not None and expected != page_count) or word_record["count"] != page_count:
            problems.append("word_page_count_mismatch")
        for index, page in enumerate(pdf):
            lines, issues = [], []
            for block in page.get_text("dict")["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    x0, y0, x1, y1 = line["bbox"]
                    if y0 > 40 and y1 < page.rect.height - 35:
                        lines.append("".join(span["text"] for span in line["spans"]))
                    if x0 < 22 or y0 < 12 or x1 > page.rect.width - 22 or y1 > page.rect.height - 12:
                        issues.append("text_near_or_outside_page_edge")
            text = normalized("".join(lines))
            page_texts.append(text)
            if index >= 3:
                for t in topics:
                    if normalized(t["number"] + t["title"]) in text:
                        if t["number"] in section_pages:
                            issues.append("duplicate_section:" + t["number"])
                        section_pages[t["number"]] = index + 1
            if len(text) < 100:
                issues.append("near_empty_page")
            image_path = output / f"word-page-{index + 1:02d}.png"
            page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False).save(image_path)
            word_images.append(image_path)
            word_pages.append({"page": index + 1, "characters": len(text), "issues": sorted(set(issues))})
            problems.extend(f"word:{index + 1}:{issue}" for issue in sorted(set(issues)))
    all_body = "".join(page_texts)
    for t in topics:
        if t["number"] not in section_pages:
            problems.append("missing_rendered_section:" + t["number"])
        for label, body in t["paragraphs"]:
            if normalized(body) not in all_body:
                problems.append("missing_rendered_paragraph:" + t["number"] + ":" + label)
        if t["table"]:
            for row in (t["table"][0], *t["table"][1]):
                for value in row:
                    if normalized(value) not in all_body:
                        problems.append("missing_rendered_cell:" + t["number"] + ":" + str(value)[:65])
    if structural["word"].get("page_map_from_render") and section_pages != structural["word"]["section_pages"]:
        problems.append("directory_fields_differ_from_actual_section_pages")
    manifest = load_json(materials / "materials_manifest.json")
    write_json(output / "word_layout.json", {
        "scope": "revised_word_directory_from_measured_render",
        "layout_source_lf_sha256": manifest["layout_source_lf_sha256"],
        "rendered_docx_raw_sha256": sha256(materials / DOCX_NAME),
        "page_count": page_count, "section_pages": section_pages,
    })
    ppt_images = sorted(ppt_dir.glob("*.png"),
                        key=lambda p: int(re.search(r"(\d+)(?=\.png$)", p.name, re.I)[1]))
    if len(ppt_images) != 20 or ppt_record["count"] != 20:
        problems.append("ppt_rendered_slide_count")
    ppt_slides = []
    for index, path in enumerate(ppt_images, 1):
        with Image.open(path) as image:
            image = image.convert("RGB")
            thumbnail = image.resize((320, 180), Image.Resampling.BILINEAR)
            ink = sum(min(thumbnail.getpixel((x, y))) < 210
                      for y in range(180) for x in range(320)) / (320 * 180)
            valid = image.size == (1600, 900) and ink >= .009 and max(ImageStat.Stat(thumbnail).stddev) > 12
            if not valid:
                problems.append(f"ppt:{index}:blank_or_wrong_size")
            ppt_slides.append({"slide": index, "size": list(image.size),
                               "dark_pixel_fraction": round(ink, 4), "nonblank": valid})
    sheets = contacts(word_images, output, "word", columns=3, rows=4, thumb_size=(260, 368))
    sheets += contacts(ppt_images, output, "ppt", columns=2, rows=2, thumb_size=(700, 394))
    result = {
        "scope": "revised_office_render_text_coverage_layout_and_pixel_checks",
        "structural": structural, "word_rendered_pages": page_count,
        "ppt_rendered_slides": len(ppt_images), "word_pages": word_pages, "ppt_slides": ppt_slides,
        "contact_sheets": sheets, "issues": problems, "automated_render_checks_passed": not problems,
        "human_visual_review_recorded": False, "ppt_render_source_equivalence": equivalent,
        "word_render_provenance": word_record, "ppt_render_provenance": ppt_record,
        "private_pdf_is_not_a_deliverable": True, "competition_ready": False,
    }
    write_json(output / "render_checks.json", result)
    print(json.dumps({"pages": page_count, "slides": len(ppt_images), "issues": problems,
                      "contact_sheets": sheets}, ensure_ascii=False, indent=2))
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--materials", type=Path, required=True)
    parser.add_argument("--word-render", type=Path)
    parser.add_argument("--ppt-render", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--ppt-render-source", type=Path)
    parser.add_argument("--draft", action="store_true",
                        help="skip document-tool hashes only; runtime and content checks stay strict")
    args = parser.parse_args(argv)
    if any((args.word_render, args.ppt_render, args.output)):
        require(all((args.word_render, args.ppt_render, args.output)), "provide all render arguments")
        result = rendered_checks(args.materials, args.word_render, args.ppt_render, args.output,
                                 check_tool_hashes=not args.draft, ppt_render_source=args.ppt_render_source)
        return 0 if result["automated_render_checks_passed"] else 2
    print(json.dumps(check_pair(args.materials, check_tool_hashes=not args.draft), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
